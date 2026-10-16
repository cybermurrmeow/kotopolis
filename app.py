from flask import jsonify
import os
import uuid
from xml.sax.saxutils import escape as _xml_escape

from datetime import date, datetime
from functools import wraps

from dotenv import load_dotenv

from flask import (
    Response,
    Flask,
    render_template,
    redirect,
    url_for,
    flash,
    request,
    abort,
    send_file
)

from flask_login import (
    LoginManager,
    login_user,
    logout_user,
    login_required,
    current_user
)

from flask_mail import Mail, Message

from itsdangerous import (
    URLSafeTimedSerializer,
    SignatureExpired,
    BadSignature
)

from sqlalchemy import or_
from sqlalchemy.exc import OperationalError, SQLAlchemyError

from werkzeug.security import (
    generate_password_hash,
    check_password_hash
)

from werkzeug.utils import secure_filename

from models import (
    db,
    User,
    Cat,
    AdoptionRequest,
    FavoriteCat,
    CatComment
)

from forms import (
    LoginForm,
    RegisterForm,
    CatForm,
    AdoptionForm,
    CommentForm,
    ForgotPasswordForm,
    ResetPasswordForm,
    ResendConfirmationForm
)



def _pdf_escape(value):

    if value is None:
        return ""

    return _xml_escape(
        str(value)
    )


# ==========================================================
# КОНФИГУРАЦИЯ
# ==========================================================

BASE_DIR = os.path.abspath(
    os.path.dirname(__file__)
)

load_dotenv(
    os.path.join(BASE_DIR, ".env")
)

from flask_wtf.csrf import CSRFProtect, CSRFError

app = Flask(__name__)



# ==========================================================
# БЕЗОПАСНЫЕ COOKIE
# ==========================================================

app.config.update(
    SESSION_COOKIE_SECURE=True,
    SESSION_COOKIE_HTTPONLY=True,
    SESSION_COOKIE_SAMESITE="Lax",

    REMEMBER_COOKIE_SECURE=True,
    REMEMBER_COOKIE_HTTPONLY=True,
    REMEMBER_COOKIE_SAMESITE="Lax",
)


# Глобальная CSRF-защита для всех изменяющих запросов.
csrf = CSRFProtect(app)


secret_key = os.environ.get("SECRET_KEY")

if not secret_key or len(secret_key) < 32:
    raise RuntimeError("SECRET_KEY must contain at least 32 characters")

app.config["SECRET_KEY"] = secret_key

app.config["SQLALCHEMY_DATABASE_URI"] = os.environ.get(
    "DATABASE_URL"
)

app.config["SQLALCHEMY_TRACK_MODIFICATIONS"] = False

app.config["UPLOAD_FOLDER"] = os.path.join(
    BASE_DIR,
    "static",
    "uploads",
    "cat_media"
)

app.config["MAX_CONTENT_LENGTH"] = 16 * 1024 * 1024


# SMTP можно будет подключить позже через .env.
# Пока отсутствие почтовых данных не ломает запуск сайта.

app.config["MAIL_SERVER"] = os.environ.get(
    "MAIL_SERVER",
    "smtp.mail.ru"
)

app.config["MAIL_PORT"] = int(
    os.environ.get(
        "MAIL_PORT",
        "465"
    )
)

app.config["MAIL_USE_SSL"] = (
    os.environ.get(
        "MAIL_USE_SSL",
        "true"
    ).lower() == "true"
)

app.config["MAIL_USE_TLS"] = (
    os.environ.get(
        "MAIL_USE_TLS",
        "false"
    ).lower() == "true"
)

app.config["MAIL_USERNAME"] = os.environ.get(
    "MAIL_USERNAME"
)

app.config["MAIL_PASSWORD"] = os.environ.get(
    "MAIL_PASSWORD"
)

app.config["MAIL_DEFAULT_SENDER"] = os.environ.get(
    "MAIL_DEFAULT_SENDER",
    app.config["MAIL_USERNAME"]
)


os.makedirs(
    app.config["UPLOAD_FOLDER"],
    exist_ok=True
)


db.init_app(app)

mail = Mail(app)

serializer = URLSafeTimedSerializer(
    app.config["SECRET_KEY"]
)


login_manager = LoginManager(app)

login_manager.login_view = "login"

login_manager.login_message = (
    "🌸 Пожалуйста, войдите в систему "
    "для доступа к этой странице"
)

login_manager.login_message_category = "warning"


# ==========================================================
# LOGIN MANAGER
# ==========================================================

# ==========================================================
# KOTOPOLIS GUEST MODE
# ==========================================================

class GuestUser:

    id = -1
    username = "Гость"
    email = ""
    role = "guest"

    email_confirmed = True
    bio = None
    created_at = None
    last_login = None

    is_authenticated = True
    is_active = True
    is_anonymous = False

    @property
    def is_admin(self):
        return False

    def get_id(self):
        return "__kotopolis_guest__"


@login_manager.user_loader
def load_user(user_id):

    if user_id == "__kotopolis_guest__":
        return GuestUser()

    try:
        return db.session.get(
            User,
            int(user_id)
        )

    except Exception as error:

        app.logger.exception(
            "Ошибка загрузки пользователя: %s",
            error
        )

        return None


# ==========================================================
# ОГРАНИЧЕНИЯ ГОСТЕВОГО РЕЖИМА
# ==========================================================

@app.before_request
def restrict_guest_mode():

    if getattr(
        current_user,
        "role",
        None
    ) != "guest":
        return None

    endpoint = request.endpoint

    if endpoint is None:
        return None

    if endpoint == "static":
        return None

    allowed_get_endpoints = {
        "index",
        "dashboard",
        "cats",
        "cat_details",
        "logout",
        "guest_login",
        "api_cats",
        "api_stats",
    }

    if (
        request.method
        in (
            "GET",
            "HEAD",
            "OPTIONS",
        )
        and
        endpoint
        in allowed_get_endpoints
    ):
        return None

    flash(
        "В гостевом режиме доступен только просмотр. "
        "Чтобы сохранять избранное, писать комментарии "
        "или отправлять заявки, войди в аккаунт.",
        "info",
    )

    return redirect(
        url_for("cats")
    )


# ==========================================================
# ПРОВЕРКА РОЛЕЙ
# ==========================================================

def staff_required(function):

    @wraps(function)
    def wrapper(*args, **kwargs):

        if not current_user.is_authenticated:
            return redirect(
                url_for("login")
            )

        if current_user.role not in (
            "admin",
            ):

            abort(403)

        return function(
            *args,
            **kwargs
        )

    return wrapper


def admin_required(function):

    @wraps(function)
    def wrapper(*args, **kwargs):

        if not current_user.is_authenticated:
            return redirect(
                url_for("login")
            )

        if current_user.role != "admin":
            abort(403)

        return function(
            *args,
            **kwargs
        )

    return wrapper


# ==========================================================
# EMAIL
# ==========================================================

def mail_is_configured():

    return bool(
        app.config.get(
            "MAIL_USERNAME"
        )
        and
        app.config.get(
            "MAIL_PASSWORD"
        )
    )


def _confirmation_resend_wait_seconds(
    user,
    cooldown_seconds=60,
):
    """
    Shared account-level cooldown for confirmation resends.

    Returns 0 when a new message may be sent, otherwise the
    remaining number of seconds.
    """

    sent_at = getattr(
        user,
        "email_code_sent_at",
        None,
    )

    if not sent_at:
        return 0

    elapsed = (
        datetime.utcnow()
        - sent_at
    ).total_seconds()

    if elapsed >= cooldown_seconds:
        return 0

    return max(
        1,
        cooldown_seconds
        - int(elapsed),
    )


def send_confirmation_email_with_cooldown(
    user,
    cooldown_seconds=60,
):
    """
    Resend wrapper that does not regenerate a code or reset
    attempts while the cooldown is active.
    """

    wait = _confirmation_resend_wait_seconds(
        user,
        cooldown_seconds=cooldown_seconds,
    )

    if wait:
        return False, wait

    sent = send_confirmation_email(
        user
    )

    return sent, 0


def send_confirmation_email(user):
    import secrets
    from datetime import datetime, timedelta
    from werkzeug.security import generate_password_hash

    if not mail_is_configured():
        app.logger.warning(
            "SMTP пока не настроен. Код подтверждения не отправлен."
        )
        return False

    code = f"{secrets.randbelow(1000000):06d}"

    user.email_code_hash = generate_password_hash(code)
    user.email_code_expires_at = datetime.utcnow() + timedelta(minutes=10)
    user.email_code_attempts = 0
    user.email_code_sent_at = datetime.utcnow()

    db.session.commit()

    message = Message(
        subject="Код подтверждения Котополиса",
        recipients=[user.email],
    )

    message.html = f"""
    <!DOCTYPE html>
    <html lang="ru">
    <head>
        <meta charset="UTF-8">
        <title>Котополис</title>
    </head>

    <body style="
        margin:0;
        padding:28px 16px;
        background:#f7f0f3;
        font-family:Arial,sans-serif;
        color:#342d31;
    ">
        <div style="
            max-width:520px;
            margin:auto;
            background:#ffffff;
            border:1px solid #eee2e8;
            border-radius:24px;
            padding:34px 28px;
            text-align:center;
        ">
            <div style="
                font-size:14px;
                font-weight:700;
                letter-spacing:.08em;
                text-transform:uppercase;
                color:#c05886;
                margin-bottom:12px;
            ">
                Котополис
            </div>

            <h1 style="
                margin:0;
                font-size:28px;
                line-height:1.15;
                color:#2f292d;
            ">
                Подтвердите почту
            </h1>

            <p style="
                color:#786c72;
                line-height:1.55;
                margin:14px 0 26px;
            ">
                Привет, <b>{user.username}</b>!
                Введите этот код в Котополисе.
            </p>

            <div style="
                display:inline-block;
                padding:18px 26px;
                border-radius:18px;
                background:#f9e7ef;
                color:#9c3e68;
                font-size:36px;
                font-weight:800;
                letter-spacing:.22em;
                margin-bottom:22px;
            ">
                {code}
            </div>

            <p style="
                color:#8c8086;
                font-size:14px;
                margin:0;
            ">
                Код действует 10 минут.
            </p>

            <p style="
                color:#aaa0a5;
                font-size:12px;
                line-height:1.5;
                margin:26px 0 0;
            ">
                Если вы не создавали аккаунт,
                просто проигнорируйте это письмо.
            </p>
        </div>
    </body>
    </html>
    """

    try:
        mail.send(message)

        app.logger.info(
            "Код подтверждения отправлен для user_id=%s",
            user.id,
        )

        return True

    except Exception as error:
        app.logger.exception(
            "Ошибка отправки кода подтверждения: %s",
            error,
        )

        return False




def send_password_reset_email(user):

    if not mail_is_configured():
        return False

    token = serializer.dumps(
        {
            "user_id": user.id,
            "email": user.email,
            "password_sig": user.password[-16:]
        },
        salt="password-reset"
    )

    reset_url = url_for(
        "reset_password",
        token=token,
        _external=True
    )

    message = Message(
        subject=(
            "🔐 Восстановление пароля - "
            "Котополис"
        ),
        recipients=[
            user.email
        ]
    )

    message.html = f"""
    <div style="
        font-family:Arial,sans-serif;
        background:#FFF0F5;
        padding:25px;
    ">

        <div style="
            max-width:600px;
            margin:auto;
            background:white;
            padding:28px;
            border-radius:18px;
        ">

            <h2 style="color:#FF1493;">
                Восстановление пароля 🔐
            </h2>

            <p>
                Ссылка запрошена для аккаунта:
                <b>{user.username}</b>
            </p>

            <p>
                Она действует в течение одного часа.
            </p>

            <p style="text-align:center;margin:30px 0;">

                <a
                    href="{reset_url}"
                    style="
                        background:#FF69B4;
                        color:white;
                        padding:13px 22px;
                        border-radius:12px;
                        text-decoration:none;
                        display:inline-block;
                    "
                >
                    Изменить пароль
                </a>

            </p>

            <p style="color:#777;font-size:14px;">
                Если вы не запрашивали смену пароля,
                просто проигнорируйте это письмо.
            </p>

        </div>

    </div>
    """

    try:

        mail.send(message)

        app.logger.info(
            "Reset-письмо отправлено "
            "для user_id=%s",
            user.id
        )

        return True

    except Exception as error:

        app.logger.exception(
            "Ошибка reset email: %s",
            error
        )

        return False


# ==========================================================
# ГЛАВНАЯ
# ==========================================================

@app.context_processor
def inject_current_year():

    return {
        "current_year":
            __import__("datetime")
            .datetime
            .now()
            .year
    }




# ==========================================================
# KOTOPOLIS AUDIT FIX: ERROR PAGES
# ==========================================================

@app.errorhandler(404)
def not_found_error(error):
    return render_template(
        "404.html"
    ), 404


@app.errorhandler(500)
def internal_server_error(error):
    db.session.rollback()

    app.logger.exception(
        "Unhandled server error",
        exc_info=error,
    )

    return render_template(
        "500.html"
    ), 500


@app.errorhandler(CSRFError)
def handle_csrf_error(error):
    return render_template(
        "csrf_error.html",
        reason=getattr(
            error,
            "description",
            "Недействительный CSRF-токен."
        )
    ), 400



# ==========================================================
# ЗАЩИТА СЕССИИ ПРИ НЕПОДТВЕРЖДЁННОМ EMAIL
# ==========================================================

@app.before_request
def enforce_confirmed_email_session():
    """
    Пользователь не должен сохранять доступ к закрытым
    страницам после того, как его email стал неподтверждённым.

    Например:
    пользователь меняет email -> email_confirmed=False ->
    текущая сессия немедленно перестаёт давать доступ.
    """

    if not current_user.is_authenticated:
        return None

    # Эти endpoint должны оставаться доступными,
    # чтобы пользователь мог подтвердить адрес,
    # повторно запросить письмо или войти заново.
    allowed_endpoints = {
        "static",
        "login",
        "logout",
        "register",
        "confirm_email",
        "resend_confirmation",
        "forgot_password",
        "reset_password",
        "demo_admin_login",
    }

    if request.endpoint in allowed_endpoints:
        return None

    if current_user.email_confirmed is not True:

        logout_user()

        flash(
            "📧 Email аккаунта не подтверждён. "
            "Подтвердите адрес, чтобы продолжить работу.",
            "warning",
        )

        return redirect(
            url_for(
                "resend_confirmation"
            )
        )

    return None



# ==========================================================
# SECURITY HEADERS
# ==========================================================

@app.before_request
def require_password_change():
    if not current_user.is_authenticated:
        return None

    if not getattr(current_user, "must_change_password", False):
        return None

    allowed = {
        "static",
        "change_password",
        "logout",
    }

    if request.endpoint in allowed:
        return None

    if request.endpoint is None:
        return None

    if request.method not in ("GET", "HEAD"):
        abort(403)

    return redirect(url_for("change_password"))


@app.after_request
def add_security_headers(response):

    # Запрещаем браузеру угадывать MIME-типы.
    response.headers[
        "X-Content-Type-Options"
    ] = "nosniff"

    # Котополис не должен встраиваться во frame/iframe.
    response.headers[
        "X-Frame-Options"
    ] = "DENY"

    # Минимизируем передачу полного URL сторонним сайтам.
    response.headers[
        "Referrer-Policy"
    ] = "strict-origin-when-cross-origin"

    # Сайт не использует эти чувствительные возможности.
    response.headers[
        "Permissions-Policy"
    ] = (
        "camera=(), "
        "microphone=(), "
        "geolocation=()"
    )

    # HTTPS уже является production-схемой сайта.
    # Браузеры игнорируют HSTS, если ответ получен не по HTTPS.
    response.headers[
        "Strict-Transport-Security"
    ] = "max-age=31536000"

    # CSP учитывает текущий Bootstrap / Font Awesome CDN,
    # inline CSS/JS существующего проекта,
    # data: favicon и внешние HTTPS-картинки.
    response.headers[
        "Content-Security-Policy"
    ] = (
        "default-src 'self'; "
        "base-uri 'self'; "
        "form-action 'self'; "
        "frame-ancestors 'none'; "
        "object-src 'none'; "

        "img-src "
        "'self' "
        "data: "
        "blob: "
        "https:; "

        "media-src "
        "'self' "
        "data: "
        "blob: "
        "https:; "

        "style-src "
        "'self' "
        "'unsafe-inline' "
        "https://cdn.jsdelivr.net "
        "https://cdnjs.cloudflare.com; "

        "script-src "
        "'self' "
        "'unsafe-inline' "
        "https://cdn.jsdelivr.net "
        "https://cdnjs.cloudflare.com; "

        "font-src "
        "'self' "
        "data: "
        "https://cdnjs.cloudflare.com "
        "https://cdn.jsdelivr.net; "

        "connect-src "
        "'self';"
    )

    # Закрытые страницы не должны кэшироваться браузером
    # или промежуточными прокси.
    try:

        if current_user.is_authenticated:

            response.headers[
                "Cache-Control"
            ] = (
                "no-store, "
                "no-cache, "
                "must-revalidate, "
                "private"
            )

            response.headers[
                "Pragma"
            ] = "no-cache"

            response.headers[
                "Expires"
            ] = "0"

    except Exception:

        pass

    return response



# ==========================================================
# КОТИКИ, КОТОРЫЕ ИЩУТ ДОМ — ДЛЯ HERO КАТАЛОГА
# ==========================================================

@app.context_processor
def inject_catalog_home_cats():

    try:

        home_cats = (
            Cat.query
            .filter(
                Cat.status
                ==
                "В приюте"
            )
            .order_by(
                Cat.id.asc()
            )
            .all()
        )

    except Exception:

        home_cats = []

    return {
        "catalog_home_cats":
            home_cats
    }


@app.route("/")
def index():

    if current_user.is_authenticated:

        return redirect(
            url_for("dashboard")
        )

    return redirect(
        url_for("login")
    )


# ==========================================================
# РЕГИСТРАЦИЯ
# ==========================================================

@app.route(
    "/register",
    methods=["GET", "POST"],
)
def register():

    if current_user.is_authenticated:
        return redirect(url_for("dashboard"))

    form = RegisterForm()

    if form.validate_on_submit():

        username = form.username.data.strip()
        email = form.email.data.strip().lower()
        password = form.password.data

        has_error = False

        existing_username = (
            User.query
            .filter(
                db.func.lower(User.username)
                ==
                username.lower()
            )
            .first()
        )

        if existing_username:
            form.username.errors.append(
                "Этот логин уже занят"
            )
            has_error = True

        if password.lower() == username.lower():
            form.password.errors.append(
                "Пароль не должен совпадать с логином"
            )
            has_error = True

        if password.lower() == email.lower():
            form.password.errors.append(
                "Пароль не должен совпадать с email"
            )
            has_error = True

        weak_passwords = {
            "password1",
            "password123",
            "qwerty123",
            "qwerty1234",
            "12345678a",
            "admin123",
            "admin1234",
        }

        if password.lower() in weak_passwords:
            form.password.errors.append(
                "Этот пароль слишком простой. Придумайте другой"
            )
            has_error = True

        if has_error:
            return render_template(
                "register.html",
                form=form,
            )

        user = User(
            username=username,
            email=email,
            password=generate_password_hash(
                password,
                method="pbkdf2:sha256",
            ),
            role="user",
            email_confirmed=False,
        )

        try:
            db.session.add(user)
            db.session.commit()

        except Exception:
            db.session.rollback()

            flash(
                "Не удалось создать аккаунт. Проверьте данные и попробуйте ещё раз.",
                "danger",
            )

            return render_template(
                "register.html",
                form=form,
            )

        sent = send_confirmation_email(user)

        ticket = serializer.dumps(
            {
                "user_id": user.id,
                "email": user.email,
            },
            salt="email-code-context",
        )

        if not sent:
            flash(
                "Аккаунт создан, но код пока не удалось отправить. "
                "Попробуйте отправить его ещё раз.",
                "warning",
            )

        return redirect(
            url_for(
                "verify_email_code",
                ticket=ticket,
            )
        )

    return render_template(
        "register.html",
        form=form,
    )




# ==========================================================
# ПОДТВЕРЖДЕНИЕ EMAIL
# ==========================================================

@app.route("/confirm_email/<token>")
def confirm_email(token):
    flash(
        "Подтверждение по ссылке больше не используется. "
        "Введите код из письма.",
        "info",
    )

    return redirect(
        url_for("login")
    )




@app.route(
    "/verify_email",
    methods=["GET", "POST"],
)
def verify_email_code():
    from datetime import datetime
    from werkzeug.security import check_password_hash

    if current_user.is_authenticated:
        return redirect(url_for("dashboard"))

    ticket = request.args.get("ticket", "").strip()

    if not ticket:
        flash(
            "Сессия подтверждения не найдена.",
            "warning",
        )
        return redirect(url_for("login"))

    try:
        payload = serializer.loads(
            ticket,
            salt="email-code-context",
            max_age=86400,
        )
    except Exception:
        flash(
            "Сессия подтверждения устарела. Войдите снова.",
            "warning",
        )
        return redirect(url_for("login"))

    try:
        user_id = int(payload.get("user_id"))
    except (TypeError, ValueError, AttributeError):
        user_id = None

    user = (
        db.session.get(User, user_id)
        if user_id
        else None
    )

    if (
        not user
        or user.email.strip().lower()
        != str(payload.get("email", "")).strip().lower()
    ):
        flash(
            "Не удалось определить аккаунт.",
            "danger",
        )
        return redirect(url_for("login"))

    if user.email_confirmed:
        flash(
            "Почта уже подтверждена. Войдите в аккаунт.",
            "info",
        )

        return redirect(url_for("login"))

    error = None

    if request.method == "POST":

        code = (
            request.form
            .get("code", "")
            .strip()
        )

        if (
            len(code) != 6
            or not code.isdigit()
        ):
            error = "Введите 6 цифр"

        elif (
            not user.email_code_hash
            or not user.email_code_expires_at
        ):
            error = "Код больше не действует. Отправьте новый."

        elif datetime.utcnow() > user.email_code_expires_at:
            error = "Код истёк. Отправьте новый."

        elif (user.email_code_attempts or 0) >= 5:
            error = "Слишком много попыток. Отправьте новый код."

        elif not check_password_hash(
            user.email_code_hash,
            code,
        ):
            user.email_code_attempts = (
                user.email_code_attempts or 0
            ) + 1

            db.session.commit()

            left = max(
                0,
                5 - user.email_code_attempts,
            )

            error = (
                f"Неверный код. Осталось попыток: {left}"
                if left
                else "Слишком много попыток. Отправьте новый код."
            )

        else:
            user.email_confirmed = True
            user.email_confirmed_at = datetime.utcnow()

            user.email_code_hash = None
            user.email_code_expires_at = None
            user.email_code_attempts = 0
            user.email_code_sent_at = None

            user.last_login = datetime.utcnow()

            db.session.commit()

            login_user(user)

            flash(
                "Почта подтверждена. Добро пожаловать в Котополис!",
                "success",
            )

            return redirect(
                url_for("dashboard")
            )

    masked_email = (
        user.email[:1]
        + "***@"
        + user.email.split("@", 1)[1]
        if "@" in user.email
        else user.email
    )

    return render_template(
        "verify_email.html",
        user=user,
        ticket=ticket,
        masked_email=masked_email,
        error=error,
    )


@app.route(
    "/resend_email_code",
    methods=["POST"],
)
def resend_email_code():
    from datetime import datetime

    ticket = request.form.get(
        "ticket",
        "",
    ).strip()

    try:
        payload = serializer.loads(
            ticket,
            salt="email-code-context",
            max_age=86400,
        )
    except Exception:
        flash(
            "Сессия подтверждения устарела.",
            "warning",
        )
        return redirect(url_for("login"))

    try:
        user_id = int(payload.get("user_id"))
    except (TypeError, ValueError, AttributeError):
        user_id = None

    user = (
        db.session.get(User, user_id)
        if user_id
        else None
    )

    if (
        not user
        or user.email.strip().lower()
        != str(
            payload.get(
                "email",
                "",
            )
        ).strip().lower()
    ):
        flash(
            "Сессия подтверждения устарела. Войдите снова.",
            "warning",
        )

        return redirect(
            url_for("login")
        )

    if user.email_confirmed:
        return redirect(url_for("login"))

    sent, wait = (
        send_confirmation_email_with_cooldown(
            user
        )
    )

    if wait:
        flash(
            f"Новый код можно отправить через {wait} сек.",
            "info",
        )

        return redirect(
            url_for(
                "verify_email_code",
                ticket=ticket,
            )
        )

    if sent:
        flash(
            "Новый код отправлен.",
            "success",
        )
    else:
        flash(
            "Не удалось отправить код. Попробуйте позже.",
            "danger",
        )

    return redirect(
        url_for(
            "verify_email_code",
            ticket=ticket,
        )
    )


@app.route(
    "/resend_confirmation",
    methods=[
        "GET",
        "POST"
    ]
)
def resend_confirmation():

    if current_user.is_authenticated:

        return redirect(
            url_for("dashboard")
        )

    form = ResendConfirmationForm()

    if form.validate_on_submit():

        username = (
            form.username.data
            .strip()
        )

        email = (
            form.email.data
            .strip()
            .lower()
        )

        user = (
            User.query
            .filter(
                db.func.lower(
                    User.username
                )
                ==
                username.lower()
            )
            .first()
        )

        if (
            user
            and
            user.email.strip().lower()
            ==
            email
        ):

            if not user.email_confirmed:

                send_confirmation_email_with_cooldown(
                    user
                )

        flash(
            "Если логин и email совпадают с аккаунтом, "
            "письмо подтверждения отправлено.",
            "info"
        )

        return redirect(
            url_for("login")
        )

    return render_template(
        "resend_confirmation.html",
        form=form
    )


# ==========================================================
# ВХОД
# ==========================================================

@app.route(
    "/login",
    methods=[
        "GET",
        "POST"
    ]
)
def login():

    if current_user.is_authenticated:

        return redirect(
            url_for("dashboard")
        )

    form = LoginForm()

    if form.validate_on_submit():

        login_value = (
            form.username_or_email.data
            .strip()
        )

        password = (
            form.password.data
        )

        user = None

        # Сначала всегда пробуем найти пользователя
        # по уникальному логину.
        user_by_username = (
            User.query
            .filter(
                db.func.lower(
                    User.username
                )
                ==
                login_value.lower()
            )
            .first()
        )

        if user_by_username:

            if check_password_hash(
                user_by_username.password,
                password
            ):

                user = user_by_username

        # Если введён email, на нём может быть
        # несколько аккаунтов.
        elif "@" in login_value:

            email_users = (
                User.query
                .filter(
                    db.func.lower(
                        User.email
                    )
                    ==
                    login_value.lower()
                )
                .all()
            )

            password_matches = [
                candidate
                for candidate in email_users
                if check_password_hash(
                    candidate.password,
                    password
                )
            ]

            if len(password_matches) == 1:

                user = password_matches[0]

            elif len(password_matches) > 1:

                flash(
                    "На эту почту зарегистрировано несколько "
                    "аккаунтов с одинаковым паролем. "
                    "Введите логин нужного аккаунта.",
                    "warning"
                )

                return render_template(
                    "login.html",
                    form=form
                )

        if not user:

            flash(
                "Неверный логин, email или пароль",
                "danger"
            )

            return render_template(
                "login.html",
                form=form
            )

        if user.email_confirmed is not True:

            sent, wait = (
                send_confirmation_email_with_cooldown(
                    user
                )
            )

            ticket = serializer.dumps(
                {
                    "user_id": user.id,
                    "email": user.email,
                },
                salt="email-code-context",
            )

            if wait:
                flash(
                    f"Новый код можно отправить через {wait} сек.",
                    "info",
                )

            elif sent:
                flash(
                    "Мы отправили новый код на вашу почту.",
                    "info",
                )

            else:
                flash(
                    "Почта ещё не подтверждена. "
                    "Код пока не удалось отправить.",
                    "warning",
                )

            return redirect(
                url_for(
                    "verify_email_code",
                    ticket=ticket,
                )
            )

        login_user(user)

        user.last_login = (
            datetime.utcnow()
        )

        db.session.commit()

        flash(
            f"Добро пожаловать, {user.username}! 🐱",
            "success"
        )

        return redirect(
            url_for("dashboard")
        )

    return render_template(
        "login.html",
        form=form
    )


# ==========================================================
# ГОСТЕВОЙ ВХОД
# ==========================================================

@app.route(
    "/guest_login",
    methods=["POST"]
)
def guest_login():

    if current_user.is_authenticated:
        return redirect(
            url_for("dashboard")
        )

    login_user(
        GuestUser(),
        remember=False
    )

    return redirect(
        url_for("dashboard")
    )


# ==========================================================
# ДЕМО-ВХОД АДМИНИСТРАТОРА
# Только для учебного проекта
# ==========================================================

@app.route("/demo_admin_login")
def demo_admin_login():

    # Публичный беспарольный вход администратора
    # отключён. Для демонстрации используется
    # обычная авторизация отдельного тестового аккаунта.
    abort(404)


# ==========================================================
# ВЫХОД
# ==========================================================

@app.route("/logout")
@login_required
def logout():

    username = current_user.username

    logout_user()

    flash(
        f"До встречи, {username}! 🐾",
        "info"
    )

    return redirect(
        url_for("login")
    )


# ==========================================================
# ВОССТАНОВЛЕНИЕ ПАРОЛЯ
# ==========================================================

@app.route(
    "/forgot_password",
    methods=[
        "GET",
        "POST"
    ]
)
def forgot_password():

    if current_user.is_authenticated:

        return redirect(
            url_for("dashboard")
        )

    form = ForgotPasswordForm()

    if form.validate_on_submit():

        username = (
            form.username.data
            .strip()
        )

        email = (
            form.email.data
            .strip()
            .lower()
        )

        user = (
            User.query
            .filter(
                db.func.lower(
                    User.username
                )
                ==
                username.lower()
            )
            .first()
        )

        if (
            user
            and
            user.email.strip().lower()
            ==
            email
        ):

            send_password_reset_email(
                user
            )

        flash(
            "Если логин и email совпадают с аккаунтом, "
            "ссылка для смены пароля отправлена.",
            "info"
        )

        return redirect(
            url_for("login")
        )

    return render_template(
        "forgot_password.html",
        form=form
    )


@app.route(
    "/reset_password/<token>",
    methods=[
        "GET",
        "POST"
    ]
)
def reset_password(token):

    if current_user.is_authenticated:

        return redirect(
            url_for("dashboard")
        )

    try:

        payload = serializer.loads(
            token,
            salt="password-reset",
            max_age=3600
        )

    except Exception:

        flash(
            "Ссылка недействительна или устарела.",
            "danger"
        )

        return redirect(
            url_for("forgot_password")
        )

    if not isinstance(
        payload,
        dict
    ):

        flash(
            "Эта ссылка была создана по старой схеме. "
            "Запросите восстановление ещё раз.",
            "warning"
        )

        return redirect(
            url_for("forgot_password")
        )

    user_id = payload.get(
        "user_id"
    )

    token_email = (
        str(
            payload.get(
                "email",
                ""
            )
        )
        .strip()
        .lower()
    )

    password_sig = str(
        payload.get(
            "password_sig",
            ""
        )
    )

    try:
        user_id = int(user_id)
    except (
        TypeError,
        ValueError
    ):
        user_id = None

    user = (
        db.session.get(
            User,
            user_id
        )
        if user_id
        else None
    )

    if (
        not user
        or
        user.email.strip().lower()
        !=
        token_email
        or
        user.password[-16:]
        !=
        password_sig
    ):

        flash(
            "Ссылка недействительна или уже была использована.",
            "danger"
        )

        return redirect(
            url_for("forgot_password")
        )

    form = ResetPasswordForm()

    if form.validate_on_submit():

        new_password = (
            form.password.data
        )

        if (
            new_password.lower()
            ==
            user.username.lower()
        ):

            form.password.errors.append(
                "Пароль не должен совпадать с логином"
            )

        weak_passwords = {
            "password1",
            "password123",
            "qwerty123",
            "qwerty1234",
            "12345678a",
            "admin123",
            "admin1234",
        }

        if (
            new_password.lower()
            in weak_passwords
        ):

            form.password.errors.append(
                "Этот пароль слишком простой. "
                "Придумайте другой"
            )

        if not form.password.errors:

            user.password = (
                generate_password_hash(
                    new_password,
                    method="pbkdf2:sha256"
                )
            )

            db.session.commit()

            flash(
                f"Пароль аккаунта {user.username} изменён 🩷",
                "success"
            )

            return redirect(
                url_for("login")
            )

    return render_template(
        "reset_password.html",
        form=form,
        token=token,
        username=user.username
    )


# ==========================================================
# DASHBOARD
# ==========================================================

@app.route("/dashboard")
@login_required
def dashboard():

    today = date.today()

    # Все котики
    cats = (
        Cat.query
        .order_by(Cat.id.desc())
        .all()
    )

    # Последние добавленные
    recent_cats = (
        Cat.query
        .order_by(Cat.id.desc())
        .limit(8)
        .all()
    )

    # Скоро в добрые руки:
    # котики на лечении и в карантине
    soon_ready = [
        cat
        for cat in cats
        if cat.status in [
            "На лечении",
            "В карантине"
        ]
    ]

    # Требуют внимания:
    # показываем только реально просроченные прививки.
    # Отсутствие медицинских данных само по себе
    # не считается проблемой.
    need_vacc = []

    for cat in cats:

        overdue_vaccine = (
            cat.next_vacc_date is not None
            and
            cat.next_vacc_date < today
        )

        if overdue_vaccine:

            cat.overdue_days = (
                today
                - cat.next_vacc_date
            ).days

            need_vacc.append(cat)

    # Медицинская статистика показывается
    # только если в базе действительно есть данные.
    vaccination_known = [
        cat
        for cat in cats
        if cat.is_vaccinated is not None
    ]

    sterilization_known = [
        cat
        for cat in cats
        if cat.is_sterilized is not None
    ]

    vaccinated_count = len([
        cat
        for cat in vaccination_known
        if cat.is_vaccinated is True
    ])

    sterilized_count = len([
        cat
        for cat in sterilization_known
        if cat.is_sterilized is True
    ])

    need_vacc_count = len(
        need_vacc
    )

    vaccination_data_available = bool(
        vaccination_known
    )

    sterilization_data_available = bool(
        sterilization_known
    )

    stats = {

        "total":
            len(cats),

        "in_shelter":
            len([
                cat
                for cat in cats
                if cat.status == "В приюте"
            ]),

        "adopted":
            len([
                cat
                for cat in cats
                if cat.status == "Усыновлён"
            ]),

        "treatment":
            len([
                cat
                for cat in cats
                if cat.status == "На лечении"
            ]),

        "quarantine":
            len([
                cat
                for cat in cats
                if cat.status == "В карантине"
            ]),

        "vaccinated_count":
            vaccinated_count,

        "sterilized_count":
            sterilized_count,

        "need_vacc_count":
            need_vacc_count,

        "vaccination_data_available":
            vaccination_data_available,

        "sterilization_data_available":
            sterilization_data_available,

        "vaccination_rate": (
            vaccinated_count
            / len(vaccination_known)
            * 100
            if vaccination_known
            else None
        ),

        "sterilization_rate": (
            sterilized_count
            / len(sterilization_known)
            * 100
            if sterilization_known
            else None
        ),

        "need_vacc_rate": (
            need_vacc_count
            / len(cats)
            * 100
            if cats
            else 0
        ),

        "moods": {}
    }

    # Настроения
    for cat in cats:

        mood = (
            cat.mood
            if cat.mood
            else "Не указано"
        )

        stats["moods"][mood] = (
            stats["moods"].get(
                mood,
                0
            )
            + 1
        )

    return render_template(
        "dashboard.html",
        cats=cats,
        stats=stats,
        recent_cats=recent_cats,
        need_vacc=need_vacc,
        soon_ready=soon_ready
    )

# ==========================================================
# КАТАЛОГ
# ==========================================================

@app.route("/cats")
@login_required
def cats():

    search = request.args.get(
        "search",
        ""
    )

    status = request.args.get(
        "status",
        ""
    )

    mood = request.args.get(
        "mood",
        ""
    )

    gender = request.args.get(
        "gender",
        ""
    )

    personality = request.args.get(
        "personality",
        ""
    )

    activity_level = request.args.get(
        "activity_level",
        ""
    )

    vocal = request.args.get(
        "vocal",
        ""
    )

    sort = request.args.get(
        "sort",
        "name"
    )

    query = Cat.query


    # Поиск
    if search:

        query = query.filter(

            Cat.name.ilike(
                f"%{search}%"
            )

            |

            Cat.nickname.ilike(
                f"%{search}%"
            )

            |

            Cat.breed.ilike(
                f"%{search}%"
            )

            |

            Cat.color.ilike(
                f"%{search}%"
            )
        )


    # Статус
    if status:

        query = query.filter(
            Cat.status == status
        )


    # Настроение
    if mood:

        query = query.filter(
            Cat.mood == mood
        )


    # Пол
    if gender:

        query = query.filter(
            Cat.gender == gender
        )


    # Характер
    if personality:

        query = query.filter(
            Cat.personality.ilike(
                f"%{personality}%"
            )
        )


    # Уровень активности
    if activity_level:

        query = query.filter(
            Cat.activity_level
            == activity_level
        )


    # Голос
    if vocal:

        query = query.filter(
            Cat.vocal == vocal
        )


    # Сортировка
    if sort == "name":

        query = query.order_by(
            Cat.name.asc()
        )

    elif sort == "name_desc":

        query = query.order_by(
            Cat.name.desc()
        )

    elif sort == "age":

        query = query.order_by(
            Cat.age.asc()
        )

    elif sort == "age_desc":

        query = query.order_by(
            Cat.age.desc()
        )

    elif sort == "newest":

        query = query.order_by(
            Cat.id.desc()
        )

    elif sort == "oldest":

        query = query.order_by(
            Cat.id.asc()
        )

    else:

        query = query.order_by(
            Cat.name.asc()
        )


    cats_list = query.all()


    # Значения для фильтров
    all_cats = Cat.query.all()

    moods = sorted({
        cat.mood
        for cat in all_cats
        if cat.mood
    })

    activity_levels = sorted({
        cat.activity_level
        for cat in all_cats
        if cat.activity_level
    })

    vocals = sorted({
        cat.vocal
        for cat in all_cats
        if cat.vocal
    })


    # Избранные текущего пользователя
    favorite_ids = []

    if current_user.is_authenticated:

        favorite_ids = [
            favorite.cat_id
            for favorite
            in FavoriteCat.query.filter_by(
                user_id=current_user.id
            ).all()
        ]


    return render_template(
        "cats.html",

        cats=cats_list,

        favorite_ids=favorite_ids,

        moods=moods,

        activity_levels=activity_levels,

        vocals=vocals,

        search=search,

        status=status,

        mood=mood,

        gender=gender,

        personality=personality,

        activity_level=activity_level,

        vocal=vocal,

        sort=sort
    )

# ==========================================================
# КАРТОЧКА КОТИКА
# ==========================================================

@app.route(
    "/cat/<int:cat_id>",
    methods=[
        "GET",
        "POST"
    ]
)
@login_required
def cat_details(cat_id):

    cat = Cat.query.get_or_404(
        cat_id
    )

    comment_form = CommentForm()

    adoption_form = AdoptionForm()

    comments = (
        CatComment.query
        .filter_by(
            cat_id=cat.id
        )
        .order_by(
            CatComment.created_at.desc()
        )
        .all()
    )

    adoption_request = (
        AdoptionRequest.query
        .filter_by(
            cat_id=cat.id,
            user_id=current_user.id
        )
        .order_by(
            AdoptionRequest.id.desc()
        )
        .first()
    )

    is_favorite = (
        FavoriteCat.query
        .filter_by(
            user_id=current_user.id,
            cat_id=cat.id
        )
        .first()
        is not None
    )

    if (
        request.method == "POST"
        and
        adoption_form.validate_on_submit()
    ):

        if cat.status != "В приюте":

            flash(
                "❌ Сейчас на этого котика "
                "нельзя подать заявку.",
                "danger"
            )

            return redirect(
                url_for(
                    "cat_details",
                    cat_id=cat.id
                )
            )

        duplicate = (
            AdoptionRequest.query
            .filter_by(
                cat_id=cat.id,
                user_id=current_user.id
            )
            .filter(
                AdoptionRequest.status.in_(
                    [
                        "Новая",
                        "На рассмотрении",
                        "Одобрена"
                    ]
                )
            )
            .first()
        )

        if duplicate:

            flash(
                "⚠️ У вас уже есть активная "
                "заявка на этого котика.",
                "warning"
            )

            return redirect(
                url_for(
                    "cat_details",
                    cat_id=cat.id
                )
            )

        adoption = AdoptionRequest(
            cat_id=cat.id,

            user_id=current_user.id,

            adopter_name=(
                adoption_form
                .adopter_name
                .data
                .strip()
            ),

            age=(
                adoption_form
                .age
                .data
            ),

            phone=(
                adoption_form
                .phone
                .data
                .strip()
            ),

            email=(
                adoption_form
                .email
                .data
                .strip()
                .lower()
            ),

            comment=(
                adoption_form.comment.data.strip()
                if adoption_form.comment.data
                else None
            ),

            request_date=(
                datetime.utcnow()
            ),

            status="Новая"
        )

        db.session.add(adoption)
        db.session.commit()

        flash(
            "💗 Заявка успешно отправлена!",
            "success"
        )

        return redirect(
            url_for(
                "cat_details",
                cat_id=cat.id
            )
        )

    return render_template(
        "cat_details.html",

        cat=cat,

        form=comment_form,

        adoption_form=adoption_form,

        comments=comments,

        adoption_request=(
            adoption_request
        ),

        is_favorite=is_favorite
    )


# ==========================================================
# КОММЕНТАРИИ
# ==========================================================





# ==========================================================
# ФОТО / GIF
# ==========================================================

def save_cat_media(form):

    from urllib.parse import urlparse

    allowed_extensions = {
        ".jpg",
        ".jpeg",
        ".png",
        ".webp",
        ".gif"
    }

    max_size = (
        10
        *
        1024
        *
        1024
    )

    # Загруженный файл имеет приоритет
    # над внешней ссылкой.
    if (
        form.photo.data
        and
        getattr(
            form.photo.data,
            "filename",
            None
        )
    ):

        original_name = secure_filename(
            form.photo.data.filename
        )

        extension = os.path.splitext(
            original_name
        )[1].lower()

        if extension not in allowed_extensions:

            raise ValueError(
                "Недопустимый формат изображения"
            )

        stream = form.photo.data.stream

        stream.seek(
            0,
            os.SEEK_END
        )

        size = stream.tell()

        stream.seek(0)

        if size <= 0:

            raise ValueError(
                "Файл изображения пустой"
            )

        if size > max_size:

            raise ValueError(
                "Размер изображения не должен "
                "превышать 10 МБ"
            )

        # Проверяем реальные байты изображения.
        # Ограничиваем также разрешение, пиксели и кадры.
        try:
            import warnings
            from PIL import Image, UnidentifiedImageError

            max_width = 10000
            max_height = 10000
            max_pixels = 40_000_000
            max_frames = 300

            with warnings.catch_warnings():
                warnings.simplefilter(
                    "error",
                    Image.DecompressionBombWarning,
                )

                with Image.open(stream) as image:

                    width = int(image.width)
                    height = int(image.height)

                    frames = int(
                        getattr(
                            image,
                            "n_frames",
                            1,
                        )
                        or 1
                    )

                    if (
                        width <= 0
                        or height <= 0
                        or width > max_width
                        or height > max_height
                        or width * height > max_pixels
                    ):
                        raise ValueError(
                            "Слишком большое разрешение изображения"
                        )

                    if frames > max_frames:
                        raise ValueError(
                            "В изображении слишком много кадров"
                        )

                    image.verify()

        except (
            Image.DecompressionBombError,
            Image.DecompressionBombWarning,
        ) as exc:

            stream.seek(0)

            raise ValueError(
                "Слишком большое разрешение изображения"
            ) from exc

        except (
            UnidentifiedImageError,
            OSError,
            ValueError,
        ) as exc:

            stream.seek(0)

            if str(exc) in {
                "Слишком большое разрешение изображения",
                "В изображении слишком много кадров",
            }:
                raise

            raise ValueError(
                "Файл повреждён или не является изображением"
            ) from exc

        finally:
            stream.seek(0)

        os.makedirs(
            app.config[
                "UPLOAD_FOLDER"
            ],
            exist_ok=True
        )

        filename = (
            f"{uuid.uuid4().hex}"
            f"{extension}"
        )

        target = os.path.join(
            app.config[
                "UPLOAD_FOLDER"
            ],
            filename
        )

        form.photo.data.save(
            target
        )

        return (
            "uploads/cat_media/"
            +
            filename
        )


    if (
        form.gif_url.data
        and
        form.gif_url.data.strip()
    ):

        url = (
            form.gif_url.data
            .strip()
        )

        parsed = urlparse(
            url
        )

        if (
            parsed.scheme
            not in {
                "http",
                "https"
            }
            or
            not parsed.netloc
        ):

            raise ValueError(
                "Ссылка на изображение должна "
                "начинаться с http:// или https://"
            )

        return url

    return None



def _remove_managed_cat_media(media_path):

    if not media_path:
        return

    if not media_path.startswith(
        "uploads/cat_media/"
    ):
        return

    if "/meowly_summer/" in media_path:
        return

    filename = os.path.basename(
        media_path
    )

    if not filename:
        return

    target = os.path.join(
        app.config[
            "UPLOAD_FOLDER"
        ],
        filename
    )

    try:

        if os.path.isfile(
            target
        ):
            os.remove(
                target
            )

    except OSError:

        app.logger.exception(
            "Failed to remove managed cat media: %s",
            media_path,
        )


def apply_cat_form(cat, form):

    old_status = (
        cat.status
    )

    new_status = (
        form.status.data
        or
        "В приюте"
    )

    allowed_statuses = {
        "В приюте",
        "Усыновлён",
        "На лечении",
        "В карантине"
    }

    if new_status not in allowed_statuses:

        raise ValueError(
            "Некорректный статус котика"
        )


    cat.name = (
        form.name.data
        .strip()
    )

    cat.nickname = (
        form.nickname.data.strip()
        if form.nickname.data
        else None
    )

    cat.age = (
        form.age.data
    )

    cat.breed = (
        form.breed.data.strip()
        if form.breed.data
        else None
    )

    cat.gender = (
        form.gender.data
        or None
    )

    cat.color = (
        form.color.data.strip()
        if form.color.data
        else None
    )

    cat.weight = (
        form.weight.data
    )

    cat.mood = (
        form.mood.data
        or None
    )

    cat.favorite_toy = (
        form.favorite_toy.data.strip()
        if form.favorite_toy.data
        else None
    )

    cat.arrival_date = (
        form.arrival_date.data
        or
        cat.arrival_date
        or
        date.today()
    )

    cat.status = (
        new_status
    )

    cat.health = (
        form.health.data.strip()
        if form.health.data
        else None
    )

    cat.description = (
        form.description.data.strip()
        if form.description.data
        else None
    )

    cat.last_vacc_date = (
        form.last_vacc_date.data
    )

    cat.next_vacc_date = (
        form.next_vacc_date.data
    )

    cat.is_vaccinated = bool(
        form.is_vaccinated.data
    )

    cat.is_sterilized = bool(
        form.is_sterilized.data
    )

    cat.personality = (
        form.personality.data
        or None
    )

    cat.activity_level = (
        form.activity_level.data
        or None
    )

    cat.kids_friendly = bool(
        form.kids_friendly.data
    )

    cat.cats_friendly = bool(
        form.cats_friendly.data
    )

    cat.dogs_friendly = bool(
        form.dogs_friendly.data
    )

    cat.favorite_place = (
        form.favorite_place.data.strip()
        if form.favorite_place.data
        else None
    )

    cat.food_preferences = (
        form.food_preferences.data.strip()
        if form.food_preferences.data
        else None
    )

    cat.vocal = (
        form.vocal.data
        or None
    )

    cat.story = (
        form.story.data.strip()
        if form.story.data
        else None
    )

    cat.adoption_recommendation = (
        form.adoption_recommendation.data.strip()
        if form.adoption_recommendation.data
        else None
    )


    if (
        cat.last_vacc_date
        and
        cat.next_vacc_date
        and
        cat.next_vacc_date
        <
        cat.last_vacc_date
    ):

        raise ValueError(
            "Дата следующей прививки "
            "не может быть раньше предыдущей"
        )


    # При переходе в статус "Усыновлён"
    # фиксируем дату и количество дней в приюте.
    if (
        new_status
        ==
        "Усыновлён"
    ):

        if (
            old_status
            !=
            "Усыновлён"
        ):

            cat.previous_status = (
                old_status
                or
                "В приюте"
            )

        cat.adoption_date = (
            form.adoption_date.data
            or
            cat.adoption_date
            or
            date.today()
        )

        if (
            cat.arrival_date
            and
            cat.adoption_date
            <
            cat.arrival_date
        ):

            raise ValueError(
                "Дата усыновления не может быть "
                "раньше даты поступления"
            )

        if (
            cat.shelter_days_frozen
            is None
        ):

            cat.shelter_days_frozen = (
                cat.days_in_shelter
            )


    else:

        # У неусыновлённого котика
        # не должно быть активной даты усыновления.
        if (
            old_status
            ==
            "Усыновлён"
        ):

            cat.previous_status = (
                None
            )

        cat.adoption_date = None

        cat.shelter_days_frozen = None


    # Save media only after all field/status/date checks.
    media = save_cat_media(
        form
    )

    if media:

        cat.photo_path = (
            media
        )


    cat.updated_at = (
        datetime.utcnow()
    )


# ==========================================================
# ДОБАВЛЕНИЕ КОТИКА
# ==========================================================

@app.route(
    "/add_cat",
    methods=[
        "GET",
        "POST"
    ]
)
@login_required
@staff_required
def add_cat():

    form = CatForm()

    if form.validate_on_submit():

        cat = Cat(
            name=form.name.data.strip(),
            arrival_date=(
                form.arrival_date.data
                or date.today()
            )
        )

        new_media_path = ""

        try:

            apply_cat_form(
                cat,
                form
            )

            new_media_path = (
                cat.photo_path
                or ""
            )

            db.session.add(
                cat
            )

            db.session.commit()

        except (
            ValueError,
            SQLAlchemyError,
        ) as exc:

            db.session.rollback()

            _remove_managed_cat_media(
                new_media_path
            )

            if isinstance(
                exc,
                ValueError,
            ):

                flash(
                    str(exc),
                    "danger"
                )

            else:

                app.logger.exception(
                    "Database error while adding cat"
                )

                flash(
                    "Не удалось сохранить котика. "
                    "Проверьте данные и попробуйте ещё раз.",
                    "danger"
                )

        else:

            flash(
                "Котик добавлен 🐱",
                "success"
            )

            return redirect(
                url_for(
                    "cat_details",
                    cat_id=cat.id
                )
            )

    return render_template(
        "add_cat.html",
        form=form
    )


# ==========================================================
# РЕДАКТИРОВАНИЕ КОТИКА
# ==========================================================

@app.route(
    "/edit_cat/<int:cat_id>",
    methods=[
        "GET",
        "POST"
    ]
)
@login_required
@staff_required
def edit_cat(cat_id):

    cat = Cat.query.get_or_404(
        cat_id
    )

    form = CatForm(
        obj=cat
    )

    if form.validate_on_submit():

        old_media_path = (
            cat.photo_path
            or ""
        )

        new_media_path = (
            old_media_path
        )

        try:

            apply_cat_form(
                cat,
                form
            )

            new_media_path = (
                cat.photo_path
                or ""
            )

            db.session.commit()

        except (
            ValueError,
            SQLAlchemyError,
        ) as exc:

            db.session.rollback()

            if (
                new_media_path
                and
                new_media_path
                !=
                old_media_path
            ):
                _remove_managed_cat_media(
                    new_media_path
                )

            if isinstance(
                exc,
                ValueError,
            ):

                flash(
                    str(exc),
                    "danger"
                )

            else:

                app.logger.exception(
                    "Database error while editing cat id=%s",
                    cat_id,
                )

                flash(
                    "Не удалось сохранить изменения. "
                    "Старые данные и изображение сохранены.",
                    "danger"
                )

        else:

            # Если изображение реально заменили,
            # удаляем старый локальный upload
            # только ПОСЛЕ успешного commit.
            if (
                old_media_path
                and
                old_media_path
                !=
                new_media_path
                and
                old_media_path.startswith(
                    "uploads/cat_media/"
                )
                and
                "/meowly_summer/"
                not in old_media_path
                and
                not old_media_path.startswith(
                    "http://"
                )
                and
                not old_media_path.startswith(
                    "https://"
                )
            ):

                old_filename = (
                    os.path.basename(
                        old_media_path
                    )
                )

                old_target = (
                    os.path.join(
                        app.config[
                            "UPLOAD_FOLDER"
                        ],
                        old_filename
                    )
                )

                try:

                    if os.path.isfile(
                        old_target
                    ):

                        os.remove(
                            old_target
                        )

                except OSError:

                    app.logger.exception(
                        "Failed to remove replaced cat media"
                    )

            flash(
                "Данные котика сохранены",
                "success"
            )

            return redirect(
                url_for(
                    "cat_details",
                    cat_id=cat.id
                )
            )

    return render_template(
        "edit_cat.html",
        form=form,
        cat=cat
    )


# ==========================================================
# УДАЛЕНИЕ КОТИКА
# ==========================================================

@app.route(
    "/delete_cat/<int:cat_id>",
    methods=[
        "POST"
    ]
)
@login_required
@admin_required
def delete_cat(cat_id):

    cat = Cat.query.get_or_404(
        cat_id
    )

    media_path = (
        cat.photo_path
        or ""
    )

    db.session.delete(
        cat
    )

    db.session.commit()

    # Удаляем только файл, загруженный
    # через админку в корневую папку cat_media.
    # Исторические GIF / WebP не трогаем.
    if (
        media_path.startswith(
            "uploads/cat_media/"
        )
        and
        "/meowly_summer/"
        not in media_path
        and
        not media_path.startswith(
            "http://"
        )
        and
        not media_path.startswith(
            "https://"
        )
    ):

        filename = os.path.basename(
            media_path
        )

        target = os.path.join(
            app.config[
                "UPLOAD_FOLDER"
            ],
            filename
        )

        try:

            if os.path.isfile(
                target
            ):

                os.remove(
                    target
                )

        except OSError:

            app.logger.exception(
                "Failed to remove cat media"
            )

    flash(
        "Котик удалён",
        "success"
    )

    return redirect(
        url_for("cats")
    )


# ==========================================================
# ИЗБРАННОЕ
# ==========================================================

@app.route(
    "/favorite/<int:cat_id>",
    methods=[
        "POST"
    ]
)
@login_required
def toggle_favorite(cat_id):

    cat = Cat.query.get_or_404(
        cat_id
    )

    is_ajax = (
        request.headers.get(
            "X-Requested-With"
        )
        ==
        "XMLHttpRequest"
    )

    favorite = (
        FavoriteCat.query
        .filter_by(
            user_id=current_user.id,
            cat_id=cat_id
        )
        .first()
    )

    if (
        cat.status == "Усыновлён"
        and
        favorite is None
    ):

        message = (
            "Этот котик уже дома и больше "
            "не доступен для добавления "
            "в избранное."
        )

        if is_ajax:
            return jsonify(
                {
                    "success": False,
                    "error": message,
                    "cat_id": cat_id
                }
            ), 409

        flash(
            message,
            "info"
        )

        return redirect(
            request.referrer
            or
            url_for("cats")
        )


    if favorite:

        db.session.delete(
            favorite
        )

        is_favorite = False

        if not is_ajax:
            flash(
                "💔 Удалено из избранного",
                "info"
            )

    else:

        favorite = FavoriteCat(
            user_id=current_user.id,
            cat_id=cat_id
        )

        db.session.add(
            favorite
        )

        is_favorite = True

        if not is_ajax:
            flash(
                "💗 Добавлено в избранное",
                "success"
            )


    db.session.commit()


    if is_ajax:

        return jsonify(
            {
                "success": True,
                "cat_id": cat_id,
                "is_favorite": is_favorite
            }
        )


    return redirect(
        request.referrer
        or
        url_for("cats")
    )


@app.route("/favorites")
@login_required
def favorites():

    favorite_rows = (
        FavoriteCat.query
        .filter_by(
            user_id=current_user.id
        )
        .order_by(
            FavoriteCat.created_at.desc()
        )
        .all()
    )

    cats_list = [
        item.cat
        for item
        in favorite_rows
    ]

    return render_template(
        "favorites.html",
        cats=cats_list
    )


# ==========================================================
# ПОДАЧА ЗАЯВКИ НА УСЫНОВЛЕНИЕ
# ==========================================================

@app.route(
    "/cat/<int:cat_id>/adopt",
    methods=[
        "GET",
        "POST"
    ]
)
@login_required
def create_adoption_request(cat_id):

    from forms import AdoptionForm

    cat = Cat.query.get_or_404(
        cat_id
    )

    # Заявки принимаются только на котиков,
    # которые сейчас готовы к усыновлению.
    if cat.status != "В приюте":

        flash(
            "На этого котика сейчас нельзя "
            "подать заявку на усыновление.",
            "warning"
        )

        return redirect(
            url_for(
                "cat_details",
                cat_id=cat.id
            )
        )

    # Один пользователь не должен создавать
    # несколько одновременно активных заявок
    # на одного и того же котика.
    active_request = (
        AdoptionRequest.query
        .filter(
            AdoptionRequest.cat_id
            ==
            cat.id,

            AdoptionRequest.user_id
            ==
            current_user.id,

            AdoptionRequest.status.in_([
                "Новая",
                "На рассмотрении",
                "Одобрена"
            ])
        )
        .order_by(
            AdoptionRequest.id.desc()
        )
        .first()
    )

    if active_request:

        flash(
            "У вас уже есть активная заявка "
            f"на {cat.name}. "
            "Её статус можно посмотреть в профиле.",
            "info"
        )

        return redirect(
            url_for("profile")
        )

    form = AdoptionForm()

    if (
        request.method
        ==
        "GET"
    ):

        form.email.data = (
            current_user.email
        )

    if form.validate_on_submit():

        # Повторная серверная проверка нужна,
        # даже если пользователь открыл форму
        # раньше изменения статуса котика.
        db.session.refresh(cat)

        if cat.status != "В приюте":

            flash(
                "Статус котика изменился. "
                "Заявку отправить уже нельзя.",
                "warning"
            )

            return redirect(
                url_for(
                    "cat_details",
                    cat_id=cat.id
                )
            )

        duplicate = (
            AdoptionRequest.query
            .filter(
                AdoptionRequest.cat_id
                ==
                cat.id,

                AdoptionRequest.user_id
                ==
                current_user.id,

                AdoptionRequest.status.in_([
                    "Новая",
                    "На рассмотрении",
                    "Одобрена"
                ])
            )
            .first()
        )

        if duplicate:

            flash(
                "Такая активная заявка уже существует.",
                "warning"
            )

            return redirect(
                url_for("profile")
            )

        adoption = AdoptionRequest(
            cat_id=cat.id,
            user_id=current_user.id,

            adopter_name=(
                form.adopter_name.data
                .strip()
            ),

            age=form.age.data,

            phone=(
                form.phone.data
                .strip()
            ),

            email=(
                form.email.data
                .strip()
                .lower()
            ),

            comment=(
                form.comment.data.strip()
                if form.comment.data
                else None
            ),

            request_date=date.today(),

            status="Новая"
        )

        db.session.add(
            adoption
        )

        db.session.commit()

        flash(
            f"Заявка на {cat.name} отправлена 🩷 "
            "Следить за её статусом можно в профиле.",
            "success"
        )

        return redirect(
            url_for("profile")
        )

    return render_template(
        "adoption_form.html",
        form=form,
        cat=cat
    )



# ==========================================================
# ЗАЯВКИ
# ==========================================================

@app.route("/adoption_requests")
@login_required
@staff_required
def adoption_requests():

    status = (
        request.args
        .get(
            "status",
            ""
        )
        .strip()
    )

    query = (
        AdoptionRequest.query
    )

    if status:

        query = query.filter(
            AdoptionRequest.status
            == status
        )

    requests_list = (
        query
        .order_by(
            AdoptionRequest.request_date.desc()
        )
        .all()
    )

    return render_template(
        "adoption_requests.html",
        requests=requests_list
    )


@app.route(
    "/adoption_request/<int:request_id>/status",
    methods=[
        "POST"
    ]
)
@login_required
@staff_required
def update_adoption_request(
    request_id
):

    adoption = (
        AdoptionRequest.query
        .get_or_404(
            request_id
        )
    )

    new_status = (
        request.form
        .get(
            "status",
            ""
        )
        .strip()
    )

    allowed = {
        "Новая",
        "На рассмотрении",
        "Одобрена",
        "Отклонена",
        "Отменена"
    }

    if new_status not in allowed:
        abort(400)

    old_status = (
        adoption.status
    )

    cat = adoption.cat

    if not cat:
        abort(404)

    # Нельзя одобрить вторую заявку,
    # если кот уже усыновлён по другой.
    if (
        new_status
        ==
        "Одобрена"
    ):

        another_approved = (
            AdoptionRequest.query
            .filter(
                AdoptionRequest.cat_id
                ==
                cat.id,

                AdoptionRequest.id
                !=
                adoption.id,

                AdoptionRequest.status
                ==
                "Одобрена"
            )
            .first()
        )

        if another_approved:

            flash(
                "У этого котика уже есть "
                "одобренная заявка.",
                "warning"
            )

            return redirect(
                url_for(
                    "adoption_requests"
                )
            )

        if (
            cat.status
            ==
            "Усыновлён"
            and
            old_status
            !=
            "Одобрена"
        ):

            flash(
                "Котик уже отмечен как усыновлённый.",
                "warning"
            )

            return redirect(
                url_for(
                    "adoption_requests"
                )
            )

        if (
            old_status
            !=
            "Одобрена"
        ):

            cat.previous_status = (
                cat.status
            )

        cat.status = (
            "Усыновлён"
        )

        cat.adoption_date = (
            date.today()
        )

        cat.shelter_days_frozen = (
            cat.days_in_shelter
        )

        adoption.status = (
            "Одобрена"
        )

        # После одобрения одной заявки
        # остальные активные заявки
        # на этого котика автоматически
        # закрываются.
        other_requests = (
            AdoptionRequest.query
            .filter(
                AdoptionRequest.cat_id
                ==
                cat.id,

                AdoptionRequest.id
                !=
                adoption.id,

                AdoptionRequest.status.in_([
                    "Новая",
                    "На рассмотрении"
                ])
            )
            .all()
        )

        for item in other_requests:

            item.status = (
                "Отклонена"
            )

        db.session.commit()

        flash(
            f"Заявка одобрена. "
            f"{cat.name} отмечен как усыновлённый 🏡",
            "success"
        )

        return redirect(
            url_for(
                "adoption_requests"
            )
        )

    # Если ранее одобренную заявку
    # возвращают в другой статус,
    # возвращаем и котика в прежнее состояние.
    if (
        old_status
        ==
        "Одобрена"
        and
        new_status
        !=
        "Одобрена"
    ):

        cat.status = (
            cat.previous_status
            or
            "В приюте"
        )

        cat.adoption_date = None

        cat.shelter_days_frozen = (
            None
        )

        cat.previous_status = (
            None
        )

    adoption.status = (
        new_status
    )

    db.session.commit()

    flash(
        "Статус заявки обновлён",
        "success"
    )

    return redirect(
        url_for(
            "adoption_requests"
        )
    )


@app.route(
    "/adoption_request/<int:request_id>/delete",
    methods=[
        "POST"
    ]
)
@login_required
@admin_required
def delete_adoption_request(
    request_id
):

    adoption = (
        AdoptionRequest.query
        .get_or_404(
            request_id
        )
    )

    if (
        adoption.status
        ==
        "Одобрена"
    ):

        flash(
            "Нельзя удалить одобренную заявку. "
            "Сначала измените её статус.",
            "warning"
        )

        return redirect(
            url_for(
                "adoption_requests"
            )
        )

    db.session.delete(
        adoption
    )

    db.session.commit()

    flash(
        "Заявка удалена",
        "success"
    )

    return redirect(
        url_for(
            "adoption_requests"
        )
    )


# ==========================================================
# ПРОФИЛЬ
# ==========================================================

@app.route(
    "/profile"
)
@login_required
def profile():

    my_requests = (
        AdoptionRequest.query
        .filter_by(
            user_id=current_user.id
        )
        .order_by(
            AdoptionRequest.request_date.desc()
        )
        .all()
    )

    favorite_count = (
        FavoriteCat.query
        .filter_by(
            user_id=current_user.id
        )
        .count()
    )

    return render_template(
        "profile.html",
        user=current_user,
        requests=my_requests,
        favorite_count=favorite_count
    )


# ==========================================================
# УПРАВЛЕНИЕ ПОЛЬЗОВАТЕЛЯМИ
# ==========================================================

@app.route("/users")
@login_required
@admin_required
def users():

    users_list = (
        User.query
        .order_by(
            User.created_at.desc()
        )
        .all()
    )

    return render_template(
        "users.html",
        users=users_list
    )


@app.route(
    "/users/create",
    methods=[
        "POST"
    ]
)
@login_required
@admin_required
def create_user():

    import re

    username = (
        request.form
        .get(
            "username",
            ""
        )
        .strip()
    )

    email = (
        request.form
        .get(
            "email",
            ""
        )
        .strip()
        .lower()
    )

    password = (
        request.form
        .get(
            "password",
            ""
        )
    )

    role = (
        request.form
        .get(
            "role",
            "user"
        )
        .strip()
    )

    errors = []

    if not (
        3
        <=
        len(username)
        <=
        30
    ):

        errors.append(
            "Логин должен быть от 3 до 30 символов."
        )

    if not re.fullmatch(
        r"[a-zA-Z0-9_а-яА-ЯёЁ-]+",
        username
    ):

        errors.append(
            "Логин может содержать только буквы, "
            "цифры, дефис и нижнее подчёркивание."
        )

    duplicate_username = (
        User.query
        .filter(
            db.func.lower(
                User.username
            )
            ==
            username.lower()
        )
        .first()
    )

    if duplicate_username:

        errors.append(
            "Пользователь с таким логином уже существует."
        )

    if (
        not email
        or
        len(email) > 120
        or
        not re.fullmatch(
            r"[^@\s]+@[^@\s]+\.[^@\s]+",
            email
        )
    ):

        errors.append(
            "Введите корректный email."
        )

    if not (
        8
        <=
        len(password)
        <=
        128
    ):

        errors.append(
            "Пароль должен быть от 8 до 128 символов."
        )

    if (
        not re.search(
            r"[A-Za-zА-Яа-яЁё]",
            password
        )
        or
        not re.search(
            r"\d",
            password
        )
    ):

        errors.append(
            "Пароль должен содержать хотя бы "
            "одну букву и одну цифру."
        )

    if (
        username
        and
        password.lower()
        ==
        username.lower()
    ):

        errors.append(
            "Пароль не должен совпадать с логином."
        )

    weak_passwords = {
        "password1",
        "password123",
        "qwerty123",
        "qwerty12345",
        "admin123",
        "admin12345",
        "12345678a",
        "123456789a"
    }

    if password.lower() in weak_passwords:

        errors.append(
            "Этот пароль слишком простой."
        )

    allowed_roles = {
        "user",
        "admin"
    }

    if role not in allowed_roles:

        errors.append(
            "Некорректная роль пользователя."
        )

    if errors:

        for error in errors:

            flash(
                error,
                "danger"
            )

        return redirect(
            url_for("users")
        )

    user = User(
        username=username,
        email=email,
        password=generate_password_hash(
            password,
            method="pbkdf2:sha256"
        ),
        role=role,
        must_change_password=True,
        email_confirmed=False,
        email_confirmed_at=None
    )

    db.session.add(
        user
    )

    db.session.commit()

    try:

        send_confirmation_email(
            user
        )

        flash(
            "Пользователь создан. "
            "На его email отправлено письмо "
            "для подтверждения аккаунта.",
            "success"
        )

    except Exception:

        app.logger.exception(
            "Failed to send confirmation email "
            "for admin-created user"
        )

        flash(
            "Пользователь создан, но письмо "
            "подтверждения отправить не удалось. "
            "Его можно запросить повторно.",
            "warning"
        )

    return redirect(
        url_for("users")
    )


@app.route(
    "/user/<int:user_id>/role",
    methods=[
        "POST"
    ]
)
@login_required
@admin_required
def update_user_role(
    user_id
):

    user = User.query.get_or_404(
        user_id
    )

    role = (
        request.form
        .get(
            "role",
            ""
        )
        .strip()
    )

    allowed_roles = {
        "user",
        "admin"
    }

    if role not in allowed_roles:
        abort(400)

    if (
        user.id
        ==
        current_user.id
        and
        role
        !=
        "admin"
    ):

        flash(
            "Нельзя снять роль администратора "
            "с самого себя.",
            "warning"
        )

        return redirect(
            url_for("users")
        )

    if (
        user.role
        ==
        "admin"
        and
        role
        !=
        "admin"
    ):

        admin_count = (
            User.query
            .filter_by(
                role="admin"
            )
            .count()
        )

        if admin_count <= 1:

            flash(
                "Нельзя снять роль у последнего "
                "администратора.",
                "warning"
            )

            return redirect(
                url_for("users")
            )

    if user.role == role:

        flash(
            "У пользователя уже установлена "
            "эта роль.",
            "info"
        )

        return redirect(
            url_for("users")
        )

    user.role = role

    db.session.commit()

    flash(
        "Роль пользователя изменена.",
        "success"
    )

    return redirect(
        url_for("users")
    )


@app.route(
    "/user/<int:user_id>/delete",
    methods=[
        "POST"
    ]
)
@login_required
@admin_required
def delete_user(
    user_id
):

    user = User.query.get_or_404(
        user_id
    )

    if (
        user.id
        ==
        current_user.id
    ):

        flash(
            "Нельзя удалить свой собственный аккаунт "
            "из панели администратора.",
            "warning"
        )

        return redirect(
            url_for("users")
        )

    if user.role == "admin":

        admin_count = (
            User.query
            .filter_by(
                role="admin"
            )
            .count()
        )

        if admin_count <= 1:

            flash(
                "Нельзя удалить последнего администратора.",
                "warning"
            )

            return redirect(
                url_for("users")
            )

    # Сохраняем историю заявок.
    # Заявка остаётся, но больше не привязана
    # к удалённому аккаунту.
    (
        AdoptionRequest.query
        .filter_by(
            user_id=user.id
        )
        .update(
            {
                "user_id":
                    None
            },
            synchronize_session=False
        )
    )

    # Удаляем избранное пользователя.
    (
        FavoriteCat.query
        .filter_by(
            user_id=user.id
        )
        .delete(
            synchronize_session=False
        )
    )

    # Если на комментарии пользователя
    # кто-то отвечал, сначала отвязываем ответы.
    comment_ids = [
        item.id
        for item
        in (
            CatComment.query
            .filter_by(
                user_id=user.id
            )
            .all()
        )
    ]

    if comment_ids:

        (
            CatComment.query
            .filter(
                CatComment.parent_id.in_(
                    comment_ids
                )
            )
            .update(
                {
                    "parent_id":
                        None
                },
                synchronize_session=False
            )
        )

        (
            CatComment.query
            .filter(
                CatComment.id.in_(
                    comment_ids
                )
            )
            .delete(
                synchronize_session=False
            )
        )

    username = user.username

    db.session.delete(
        user
    )

    db.session.commit()

    flash(
        f"Пользователь {username} удалён.",
        "success"
    )

    return redirect(
        url_for("users")
    )


# ==========================================================
# ПРОСТОЙ CSV EXPORT
# Позже восстановим исходные PDF/Excel функции.
# ==========================================================

@app.route("/exports")
@login_required
@staff_required
def exports_center():

    return render_template(
        "exports.html"
    )


@app.route("/export_cats_csv")
@login_required
@staff_required
def export_cats_csv():

    import csv
    import io

    cats_list = (
        Cat.query
        .order_by(
            Cat.id.asc()
        )
        .all()
    )

    output = io.StringIO(
        newline=""
    )

    # Excel на Windows нормально распознаёт
    # UTF-8 с BOM.
    output.write(
        "\ufeff"
    )

    writer = csv.writer(
        output,
        delimiter=";",
        quoting=csv.QUOTE_MINIMAL
    )

    writer.writerow(
        [
            "ID",
            "Имя",
            "Прозвище",
            "Возраст",
            "Порода",
            "Пол",
            "Окрас",
            "Вес, кг",
            "Статус",
            "Дата поступления",
            "Дата усыновления",
            "Вакцинирован",
            "Стерилизован",
            "Последняя прививка",
            "Следующая прививка",
            "Характер",
            "Активность",
            "С детьми",
            "С кошками",
            "С собаками",
            "Здоровье",
            "Описание"
        ]
    )

    for cat in cats_list:

        writer.writerow(
            [
                cat.id,
                cat.name or "",
                cat.nickname or "",
                (
                    cat.age
                    if cat.age is not None
                    else ""
                ),
                cat.breed or "",
                cat.gender or "",
                cat.color or "",
                (
                    cat.weight
                    if cat.weight is not None
                    else ""
                ),
                cat.status or "",
                (
                    cat.arrival_date.strftime(
                        "%d.%m.%Y"
                    )
                    if cat.arrival_date
                    else ""
                ),
                (
                    cat.adoption_date.strftime(
                        "%d.%m.%Y"
                    )
                    if cat.adoption_date
                    else ""
                ),
                (
                    "Да"
                    if cat.is_vaccinated
                    else "Нет"
                ),
                (
                    "Да"
                    if cat.is_sterilized
                    else "Нет"
                ),
                (
                    cat.last_vacc_date.strftime(
                        "%d.%m.%Y"
                    )
                    if cat.last_vacc_date
                    else ""
                ),
                (
                    cat.next_vacc_date.strftime(
                        "%d.%m.%Y"
                    )
                    if cat.next_vacc_date
                    else ""
                ),
                cat.personality or "",
                cat.activity_level or "",
                (
                    "Да"
                    if cat.kids_friendly
                    else "Нет"
                ),
                (
                    "Да"
                    if cat.cats_friendly
                    else "Нет"
                ),
                (
                    "Да"
                    if cat.dogs_friendly
                    else "Нет"
                ),
                cat.health or "",
                cat.description or ""
            ]
        )

    data = output.getvalue()

    return Response(
        data,
        mimetype="text/csv",
        headers={
            "Content-Disposition":
                "attachment; "
                "filename=kotopolis_cats.csv"
        }
    )


@app.route("/export_cats_excel")
@login_required
@staff_required
def export_cats_excel():

    import io

    from openpyxl import Workbook

    from openpyxl.styles import (
        Font,
        Alignment
    )

    from openpyxl.utils import (
        get_column_letter
    )


    cats_list = (
        Cat.query
        .order_by(
            Cat.id.asc()
        )
        .all()
    )


    wb = Workbook()

    ws = wb.active

    ws.title = (
        "Котики"
    )


    headers = [
        "ID",
        "Имя",
        "Прозвище",
        "Возраст",
        "Порода",
        "Пол",
        "Окрас",
        "Вес, кг",
        "Статус",
        "Дата поступления",
        "Дата усыновления",
        "Вакцинирован",
        "Стерилизован",
        "Последняя прививка",
        "Следующая прививка",
        "Характер",
        "Активность",
        "С детьми",
        "С кошками",
        "С собаками",
        "Здоровье",
        "Описание"
    ]


    ws.append(
        headers
    )


    for cell in ws[1]:

        cell.font = Font(
            bold=True
        )

        cell.alignment = Alignment(
            horizontal="center",
            vertical="center"
        )


    for cat in cats_list:

        ws.append(
            [
                cat.id,
                cat.name or "",
                cat.nickname or "",
                cat.age,
                cat.breed or "",
                cat.gender or "",
                cat.color or "",
                cat.weight,
                cat.status or "",
                cat.arrival_date,
                cat.adoption_date,
                (
                    "Да"
                    if cat.is_vaccinated
                    else "Нет"
                ),
                (
                    "Да"
                    if cat.is_sterilized
                    else "Нет"
                ),
                cat.last_vacc_date,
                cat.next_vacc_date,
                cat.personality or "",
                cat.activity_level or "",
                (
                    "Да"
                    if cat.kids_friendly
                    else "Нет"
                ),
                (
                    "Да"
                    if cat.cats_friendly
                    else "Нет"
                ),
                (
                    "Да"
                    if cat.dogs_friendly
                    else "Нет"
                ),
                cat.health or "",
                cat.description or ""
            ]
        )


    for row in ws.iter_rows(
        min_row=2
    ):

        for index in [
            10,
            11,
            14,
            15
        ]:

            cell = row[
                index - 1
            ]

            if cell.value:

                cell.number_format = (
                    "DD.MM.YYYY"
                )


    ws.freeze_panes = (
        "A2"
    )

    ws.auto_filter.ref = (
        ws.dimensions
    )


    widths = {
        1: 7,
        2: 16,
        3: 20,
        4: 10,
        5: 18,
        6: 12,
        7: 18,
        8: 11,
        9: 18,
        10: 17,
        11: 17,
        12: 15,
        13: 15,
        14: 18,
        15: 18,
        16: 24,
        17: 16,
        18: 12,
        19: 12,
        20: 12,
        21: 35,
        22: 45
    }


    for index, width in widths.items():

        ws.column_dimensions[
            get_column_letter(
                index
            )
        ].width = width


    # Вторая вкладка со статистикой.
    stats = wb.create_sheet(
        "Статистика"
    )


    status_counts = {}

    for cat in cats_list:

        status = (
            cat.status
            or
            "Не указан"
        )

        status_counts[
            status
        ] = (
            status_counts.get(
                status,
                0
            )
            +
            1
        )


    stats.append(
        [
            "Показатель",
            "Значение"
        ]
    )

    stats.append(
        [
            "Всего котиков",
            len(
                cats_list
            )
        ]
    )

    stats.append(
        [
            "Вакцинированы",
            sum(
                1
                for cat
                in cats_list
                if cat.is_vaccinated
            )
        ]
    )

    stats.append(
        [
            "Стерилизованы",
            sum(
                1
                for cat
                in cats_list
                if cat.is_sterilized
            )
        ]
    )


    for status in sorted(
        status_counts
    ):

        stats.append(
            [
                status,
                status_counts[
                    status
                ]
            ]
        )


    for cell in stats[1]:

        cell.font = Font(
            bold=True
        )


    stats.column_dimensions[
        "A"
    ].width = 28

    stats.column_dimensions[
        "B"
    ].width = 14


    buffer = io.BytesIO()

    wb.save(
        buffer
    )

    buffer.seek(0)


    return send_file(
        buffer,
        mimetype=(
            "application/"
            "vnd.openxmlformats-"
            "officedocument."
            "spreadsheetml.sheet"
        ),
        as_attachment=True,
        download_name=(
            "kotopolis_cats.xlsx"
        )
    )


@app.route("/export_cats_pdf")
@login_required
@staff_required
def export_cats_pdf():

    import io
    import os

    from reportlab.lib import colors

    from reportlab.lib.enums import (
        TA_CENTER
    )

    from reportlab.lib.pagesizes import (
        A4
    )

    from reportlab.lib.styles import (
        getSampleStyleSheet,
        ParagraphStyle
    )

    from reportlab.lib.units import (
        mm
    )

    from reportlab.pdfbase import (
        pdfmetrics
    )

    from reportlab.pdfbase.ttfonts import (
        TTFont
    )

    from reportlab.platypus import (
        SimpleDocTemplate,
        Paragraph,
        Spacer,
        Table,
        TableStyle,
        PageBreak
    )


    cats_list = (
        Cat.query
        .order_by(
            Cat.id.asc()
        )
        .all()
    )


    buffer = io.BytesIO()


    regular_candidates = [
        "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf",
        "/usr/share/fonts/truetype/liberation2/LiberationSans-Regular.ttf",
        "C:/Windows/Fonts/arial.ttf",
        "C:/Windows/Fonts/calibri.ttf",
    ]

    bold_candidates = [
        "/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf",
        "/usr/share/fonts/truetype/liberation2/LiberationSans-Bold.ttf",
        "C:/Windows/Fonts/arialbd.ttf",
        "C:/Windows/Fonts/calibrib.ttf",
    ]

    font_path = next(
        (
            path
            for path in regular_candidates
            if os.path.isfile(path)
        ),
        None,
    )

    bold_font_path = next(
        (
            path
            for path in bold_candidates
            if os.path.isfile(path)
        ),
        None,
    )

    if not font_path:

        raise RuntimeError(
            "Не найден шрифт с поддержкой кириллицы"
        )

    pdfmetrics.registerFont(
        TTFont(
            "DejaVu",
            font_path
        )
    )

    if bold_font_path:

        pdfmetrics.registerFont(
            TTFont(
                "DejaVuBold",
                bold_font_path
            )
        )

        bold_font = "DejaVuBold"

    else:

        bold_font = "DejaVu"


    doc = SimpleDocTemplate(
        buffer,
        pagesize=A4,
        rightMargin=14 * mm,
        leftMargin=14 * mm,
        topMargin=14 * mm,
        bottomMargin=14 * mm,
        title="Котополис"
    )


    styles = getSampleStyleSheet()


    title_style = ParagraphStyle(
        "KotopolisTitle",
        parent=styles[
            "Title"
        ],
        fontName=bold_font,
        fontSize=18,
        leading=22,
        alignment=TA_CENTER,
        spaceAfter=8 * mm
    )


    heading_style = ParagraphStyle(
        "KotopolisHeading",
        parent=styles[
            "Heading2"
        ],
        fontName=bold_font,
        fontSize=11,
        leading=14
    )


    normal_style = ParagraphStyle(
        "KotopolisNormal",
        parent=styles[
            "BodyText"
        ],
        fontName="DejaVu",
        fontSize=8.5,
        leading=12
    )


    story = []


    story.append(
        Paragraph(
            "Котополис",
            title_style
        )
    )

    story.append(
        Paragraph(
            (
                "Отчёт по котикам приюта. "
                f"Всего записей: {len(cats_list)}"
            ),
            normal_style
        )
    )

    story.append(
        Spacer(
            1,
            5 * mm
        )
    )


    stats_data = [
        [
            "Статус",
            "Количество"
        ]
    ]


    statuses = [
        "В приюте",
        "Усыновлён",
        "На лечении",
        "В карантине"
    ]


    for status in statuses:

        stats_data.append(
            [
                status,
                sum(
                    1
                    for cat
                    in cats_list
                    if cat.status
                    ==
                    status
                )
            ]
        )


    stats_table = Table(
        stats_data,
        colWidths=[
            65 * mm,
            35 * mm
        ]
    )


    stats_table.setStyle(
        TableStyle(
            [
                (
                    "FONTNAME",
                    (
                        0,
                        0
                    ),
                    (
                        -1,
                        -1
                    ),
                    "DejaVu"
                ),
                (
                    "FONTNAME",
                    (
                        0,
                        0
                    ),
                    (
                        -1,
                        0
                    ),
                    bold_font
                ),
                (
                    "GRID",
                    (
                        0,
                        0
                    ),
                    (
                        -1,
                        -1
                    ),
                    0.4,
                    colors.grey
                ),
                (
                    "ALIGN",
                    (
                        1,
                        0
                    ),
                    (
                        1,
                        -1
                    ),
                    "CENTER"
                ),
                (
                    "BACKGROUND",
                    (
                        0,
                        0
                    ),
                    (
                        -1,
                        0
                    ),
                    colors.whitesmoke
                ),
                (
                    "BOTTOMPADDING",
                    (
                        0,
                        0
                    ),
                    (
                        -1,
                        -1
                    ),
                    5
                ),
                (
                    "TOPPADDING",
                    (
                        0,
                        0
                    ),
                    (
                        -1,
                        -1
                    ),
                    5
                )
            ]
        )
    )


    story.append(
        stats_table
    )

    story.append(
        PageBreak()
    )


    for index, cat in enumerate(
        cats_list,
        start=1
    ):

        story.append(
            Paragraph(
                (
                    f"{index}. "
                    f"{_pdf_escape(cat.name)}"
                ),
                heading_style
            )
        )


        data = [
            [
                "Статус",
                cat.status or "Не указан"
            ],
            [
                "Возраст",
                (
                    str(
                        cat.age
                    )
                    if cat.age is not None
                    else "Не указан"
                )
            ],
            [
                "Порода",
                cat.breed or "Не указана"
            ],
            [
                "Пол",
                cat.gender or "Не указан"
            ],
            [
                "Окрас",
                cat.color or "Не указан"
            ],
            [
                "Вес",
                (
                    f"{cat.weight} кг"
                    if cat.weight is not None
                    else "Не указан"
                )
            ],
            [
                "Поступление",
                (
                    cat.arrival_date.strftime(
                        "%d.%m.%Y"
                    )
                    if cat.arrival_date
                    else "Не указано"
                )
            ],
            [
                "Вакцинация",
                (
                    "Да"
                    if cat.is_vaccinated
                    else "Нет"
                )
            ],
            [
                "Стерилизация",
                (
                    "Да"
                    if cat.is_sterilized
                    else "Нет"
                )
            ]
        ]


        table = Table(
            data,
            colWidths=[
                42 * mm,
                118 * mm
            ]
        )


        table.setStyle(
            TableStyle(
                [
                    (
                        "FONTNAME",
                        (
                            0,
                            0
                        ),
                        (
                            -1,
                            -1
                        ),
                        "DejaVu"
                    ),
                    (
                        "FONTNAME",
                        (
                            0,
                            0
                        ),
                        (
                            0,
                            -1
                        ),
                        bold_font
                    ),
                    (
                        "VALIGN",
                        (
                            0,
                            0
                        ),
                        (
                            -1,
                            -1
                        ),
                        "TOP"
                    ),
                    (
                        "GRID",
                        (
                            0,
                            0
                        ),
                        (
                            -1,
                            -1
                        ),
                        0.25,
                        colors.lightgrey
                    ),
                    (
                        "BOTTOMPADDING",
                        (
                            0,
                            0
                        ),
                        (
                            -1,
                            -1
                        ),
                        4
                    ),
                    (
                        "TOPPADDING",
                        (
                            0,
                            0
                        ),
                        (
                            -1,
                            -1
                        ),
                        4
                    )
                ]
            )
        )


        story.append(
            table
        )


        if cat.health:

            story.append(
                Spacer(
                    1,
                    2 * mm
                )
            )

            story.append(
                Paragraph(
                    (
                        "<b>Здоровье:</b> "
                        +
                        _pdf_escape(cat.health)
                    ),
                    normal_style
                )
            )


        if cat.description:

            story.append(
                Paragraph(
                    (
                        "<b>Описание:</b> "
                        +
                        _pdf_escape(cat.description)
                    ),
                    normal_style
                )
            )


        story.append(
            Spacer(
                1,
                6 * mm
            )
        )


    doc.build(
        story
    )


    buffer.seek(0)


    return send_file(
        buffer,
        mimetype="application/pdf",
        as_attachment=True,
        download_name=(
            "kotopolis_cats.pdf"
        )
    )


@app.route(
    "/edit_profile",
    methods=[
        "GET",
        "POST"
    ]
)
@login_required
def edit_profile():

    errors = {}

    if request.method == "POST":

        # FIX PACK 5 / FRONT-01:
        # Compact profile form sends only bio + CSRF.
        is_bio_only = (
            "bio" in request.form
            and "username" not in request.form
            and "email" not in request.form
        )

        if is_bio_only:
            bio = (
                request.form
                .get(
                    "bio",
                    ""
                )
                .strip()
            )

            if len(bio) > 600:
                flash(
                    "Описание не должно превышать 600 символов",
                    "danger",
                )
                return redirect(
                    url_for("profile")
                )

            current_user.bio = bio

            try:
                db.session.commit()
            except Exception:
                db.session.rollback()
                app.logger.exception(
                    "Failed to save profile bio for user_id=%s",
                    current_user.id,
                )
                flash(
                    "Не удалось сохранить описание. Попробуйте ещё раз.",
                    "danger",
                )
            else:
                flash(
                    "Описание сохранено",
                    "success",
                )

            return redirect(
                url_for("profile")
            )

        username = (
            request.form
            .get(
                "username",
                ""
            )
            .strip()
        )

        email = (
            request.form
            .get(
                "email",
                ""
            )
            .strip()
            .lower()
        )

        bio = (
            request.form
            .get(
                "bio",
                ""
            )
            .strip()
        )

        # Логин
        if not username:

            errors["username"] = (
                "Введите логин"
            )

        elif (
            len(username) < 3
            or
            len(username) > 30
        ):

            errors["username"] = (
                "Логин должен быть от 3 до 30 символов"
            )

        else:

            import re

            if not re.fullmatch(
                r"[a-zA-Z0-9_а-яА-ЯёЁ-]+",
                username
            ):

                errors["username"] = (
                    "Используйте только буквы, цифры, _ и -"
                )

        # Проверяем уникальность логина
        if "username" not in errors:

            existing_username = (
                User.query
                .filter(
                    db.func.lower(
                        User.username
                    )
                    ==
                    username.lower(),
                    User.id
                    !=
                    current_user.id
                )
                .first()
            )

            if existing_username:

                errors["username"] = (
                    "Этот логин уже занят"
                )

        # Email
        if not email:

            errors["email"] = (
                "Введите email"
            )

        elif len(email) > 120:

            errors["email"] = (
                "Email не должен превышать 120 символов"
            )

        else:

            import re

            if not re.fullmatch(
                r"[^\s@]+@[^\s@]+\.[^\s@]+",
                email
            ):

                errors["email"] = (
                    "Введите корректный email"
                )

        # Bio
        if len(bio) > 600:

            errors["bio"] = (
                "Описание не должно превышать 600 символов"
            )

        if errors:

            return render_template(
                "edit_profile.html",
                user=current_user,
                errors=errors,
                username=username,
                email=email,
                bio=bio
            )

        old_email = (
            current_user.email
            .strip()
            .lower()
        )

        email_changed = (
            old_email
            !=
            email
        )

        current_user.username = (
            username
        )

        current_user.email = (
            email
        )

        current_user.bio = (
            bio
            or None
        )

        if email_changed:

            current_user.email_confirmed = (
                False
            )

            current_user.email_confirmed_at = (
                None
            )

        db.session.commit()

        if email_changed:

            sent = send_confirmation_email(
                current_user
            )

            ticket = serializer.dumps(
                {
                    "user_id": current_user.id,
                    "email": current_user.email,
                },
                salt="email-code-context",
            )

            logout_user()

            if sent:

                flash(
                    "Профиль сохранён 🩷 "
                    "На новую почту отправлен код "
                    "подтверждения.",
                    "success"
                )

            else:

                flash(
                    "Профиль сохранён, но код пока "
                    "не удалось отправить. "
                    "Попробуйте отправить его ещё раз.",
                    "warning"
                )

            return redirect(
                url_for(
                    "verify_email_code",
                    ticket=ticket,
                )
            )

        flash(
            "Профиль сохранён 🩷",
            "success"
        )

        return redirect(
            url_for("profile")
        )

    return render_template(
        "edit_profile.html",
        user=current_user,
        errors={},
        username=current_user.username,
        email=current_user.email,
        bio=current_user.bio or ""
    )


@app.route(
    "/change_password",
    methods=[
        "GET",
        "POST"
    ]
)
@login_required
def change_password():

    errors = {}

    if request.method == "POST":

        old_password = (
            request.form
            .get(
                "old_password",
                ""
            )
        )

        new_password = (
            request.form
            .get(
                "new_password",
                ""
            )
        )

        confirm_password = (
            request.form
            .get(
                "confirm_password",
                ""
            )
        )

        if not old_password:

            errors["old_password"] = (
                "Введите текущий пароль"
            )

        elif not check_password_hash(
            current_user.password,
            old_password
        ):

            errors["old_password"] = (
                "Текущий пароль указан неверно"
            )

        if not new_password:

            errors["new_password"] = (
                "Введите новый пароль"
            )

        elif (
            len(new_password) < 8
            or
            len(new_password) > 128
        ):

            errors["new_password"] = (
                "Пароль должен быть от 8 до 128 символов"
            )

        elif (
            not any(
                ch.isalpha()
                for ch in new_password
            )
            or
            not any(
                ch.isdigit()
                for ch in new_password
            )
        ):

            errors["new_password"] = (
                "Добавьте хотя бы одну букву и одну цифру"
            )

        elif (
            new_password.lower()
            ==
            current_user.username.lower()
        ):

            errors["new_password"] = (
                "Пароль не должен совпадать с логином"
            )

        weak_passwords = {
            "password1",
            "password123",
            "qwerty123",
            "qwerty1234",
            "12345678a",
            "admin123",
            "admin1234",
        }

        if (
            new_password
            and
            new_password.lower()
            in weak_passwords
        ):

            errors["new_password"] = (
                "Этот пароль слишком простой. "
                "Придумайте другой"
            )

        if (
            old_password
            and
            new_password
            and
            check_password_hash(
                current_user.password,
                new_password
            )
        ):

            errors["new_password"] = (
                "Новый пароль должен отличаться от текущего"
            )

        if (
            new_password
            !=
            confirm_password
        ):

            errors["confirm_password"] = (
                "Пароли не совпадают"
            )

        if errors:

            return render_template(
                "change_password.html",
                errors=errors
            )

        current_user.password = (
            generate_password_hash(
                new_password,
                method="pbkdf2:sha256"
            )
        )

        current_user.must_change_password = False
        db.session.commit()

        flash(
            "Пароль успешно изменён 🔐",
            "success"
        )

        return redirect(
            url_for("profile")
        )

    return render_template(
        "change_password.html",
        errors={}
    )


@app.route(
    "/cancel_my_request/<int:request_id>",
    methods=[
        "POST"
    ]
)
@login_required
def cancel_my_request(
    request_id
):

    adoption = (
        AdoptionRequest.query
        .get_or_404(
            request_id
        )
    )

    if (
        adoption.user_id
        !=
        current_user.id
    ):

        abort(403)

    current_status = (
        adoption.status
        or ""
    ).strip()

    cancellable = {
        "Новая",
        "На рассмотрении"
    }

    if current_status not in cancellable:

        flash(
            "Эту заявку уже нельзя отменить.",
            "warning"
        )

        return redirect(
            url_for("profile")
        )

    adoption.status = (
        "Отменена"
    )

    db.session.commit()

    flash(
        "Заявка отменена",
        "success"
    )

    return redirect(
        url_for("profile")
    )




@app.route(
    "/api/cats"
)
@login_required
def api_cats():

    from flask import jsonify

    cats_list = (
        Cat.query
        .order_by(
            Cat.id.asc()
        )
        .all()
    )


    data = []


    for cat in cats_list:

        data.append(
            {
                "id":
                    cat.id,

                "name":
                    cat.name,

                "nickname":
                    cat.nickname,

                "age":
                    cat.age,

                "age_category":
                    cat.age_category(),

                "breed":
                    cat.breed,

                "gender":
                    cat.gender,

                "color":
                    cat.color,

                "weight":
                    cat.weight,

                "mood":
                    cat.mood,

                "status":
                    cat.status,

                "health":
                    cat.health,

                "description":
                    cat.description,

                "personality":
                    cat.personality,

                "activity_level":
                    cat.activity_level,

                "kids_friendly":
                    bool(
                        cat.kids_friendly
                    ),

                "cats_friendly":
                    bool(
                        cat.cats_friendly
                    ),

                "dogs_friendly":
                    bool(
                        cat.dogs_friendly
                    ),

                "favorite_toy":
                    cat.favorite_toy,

                "favorite_place":
                    cat.favorite_place,

                "food_preferences":
                    cat.food_preferences,

                "vocal":
                    cat.vocal,

                "story":
                    cat.story,

                "adoption_recommendation":
                    cat.adoption_recommendation,

                "photo_path":
                    cat.photo_path,

                "arrival_date":
                    (
                        cat.arrival_date.isoformat()
                        if cat.arrival_date
                        else None
                    ),

                "adoption_date":
                    (
                        cat.adoption_date.isoformat()
                        if cat.adoption_date
                        else None
                    ),

                "days_in_shelter":
                    cat.days_in_shelter,

                "is_vaccinated":
                    bool(
                        cat.is_vaccinated
                    ),

                "vaccination_status":
                    cat.vacc_status(),

                "last_vacc_date":
                    (
                        cat.last_vacc_date.isoformat()
                        if cat.last_vacc_date
                        else None
                    ),

                "next_vacc_date":
                    (
                        cat.next_vacc_date.isoformat()
                        if cat.next_vacc_date
                        else None
                    ),

                "is_sterilized":
                    bool(
                        cat.is_sterilized
                    ),

                "created_at":
                    (
                        cat.created_at.isoformat()
                        if cat.created_at
                        else None
                    ),

                "updated_at":
                    (
                        cat.updated_at.isoformat()
                        if cat.updated_at
                        else None
                    )
            }
        )


    response = jsonify(
        {
            "count":
                len(
                    data
                ),

            "cats":
                data
        }
    )

    response.headers[
        "Cache-Control"
    ] = (
        "no-store"
    )

    return response


@app.route(
    "/api/stats"
)
@login_required
def api_stats():

    from flask import jsonify

    cats_list = (
        Cat.query.all()
    )


    total = len(
        cats_list
    )


    by_status = {}

    for cat in cats_list:

        status = (
            cat.status
            or
            "Не указан"
        )

        by_status[
            status
        ] = (
            by_status.get(
                status,
                0
            )
            +
            1
        )


    vaccinated = sum(
        1
        for cat
        in cats_list
        if cat.is_vaccinated
    )


    sterilized = sum(
        1
        for cat
        in cats_list
        if cat.is_sterilized
    )


    attention = sum(
        1
        for cat
        in cats_list
        if (
            cat.next_vacc_date
            and
            cat.next_vacc_date
            <
            date.today()
        )
    )


    response = jsonify(
        {
            "total":
                total,

            "by_status":
                by_status,

            "in_shelter":
                by_status.get(
                    "В приюте",
                    0
                ),

            "adopted":
                by_status.get(
                    "Усыновлён",
                    0
                ),

            "treatment":
                by_status.get(
                    "На лечении",
                    0
                ),

            "quarantine":
                by_status.get(
                    "В карантине",
                    0
                ),

            "vaccinated":
                vaccinated,

            "vaccinated_percent":
                (
                    round(
                        vaccinated
                        /
                        total
                        *
                        100,
                        1
                    )
                    if total
                    else 0
                ),

            "sterilized":
                sterilized,

            "sterilized_percent":
                (
                    round(
                        sterilized
                        /
                        total
                        *
                        100,
                        1
                    )
                    if total
                    else 0
                ),

            "vaccination_attention":
                attention
        }
    )


    response.headers[
        "Cache-Control"
    ] = (
        "no-store"
    )

    return response


@app.route(
    "/export_excel"
)
@login_required
@admin_required
def export_excel():

    from utils import export_to_excel

    cats_list = (
        Cat.query
        .order_by(
            Cat.id.asc()
        )
        .all()
    )

    excel_file = (
        export_to_excel(
            cats_list
        )
    )

    return send_file(
        excel_file,
        as_attachment=True,
        download_name=(
            "kotopolis_cats_"
            + date.today().isoformat()
            + ".xlsx"
        ),
        mimetype=(
            "application/"
            "vnd.openxmlformats-officedocument."
            "spreadsheetml.sheet"
        )
    )


@app.route(
    "/report"
)
@login_required
@admin_required
def report():

    from utils import (
        calculate_stats,
        generate_pdf_report,
    )

    cats_list = (
        Cat.query
        .order_by(
            Cat.id.desc()
        )
        .all()
    )

    stats = (
        calculate_stats(
            cats_list
        )
    )

    pdf_file = (
        generate_pdf_report(
            cats_list,
            stats
        )
    )

    return send_file(
        pdf_file,
        as_attachment=True,
        download_name=(
            "kotopolis_report_"
            + date.today().isoformat()
            + ".pdf"
        ),
        mimetype="application/pdf"
    )


# ==========================================================
# END RESTORED OLD KOTOPOLIS EXPORT ROUTES
# ==========================================================



# ===== DAILY STORY START =====
from daily_story import register_daily_story
register_daily_story(app)
# ===== DAILY STORY END =====

@app.route("/favicon.ico")
def favicon():
    return app.send_static_file("icons/favicon.ico")



if __name__ == "__main__":

    port = int(
        os.environ.get(
            "PORT",
            5000
        )
    )

    app.run(
        host="0.0.0.0",
        port=port,
        debug=False
    )


# FIX PACK 6 V2: backend audit fixes 2026-10-04


# FIX PACK 7: confirmed security audit fixes 2026-10-04
