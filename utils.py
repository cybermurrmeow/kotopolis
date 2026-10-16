from collections import Counter
from datetime import date, datetime
from io import BytesIO
import os
from xml.sax.saxutils import escape as _xml_escape

from openpyxl import Workbook
from openpyxl.styles import (
    Alignment,
    Font,
    PatternFill,
    Border,
    Side,
)

from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import (
    getSampleStyleSheet,
    ParagraphStyle,
)
from reportlab.lib.units import mm
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import (
    SimpleDocTemplate,
    Paragraph,
    Spacer,
    Table,
    TableStyle,
    PageBreak,
)



def _pdf_escape(value):

    if value is None:
        return ""

    return _xml_escape(
        str(value)
    )


# ==========================================================
# СТАТИСТИКА
# ==========================================================

def calculate_stats(cats):

    cats = list(cats)

    total = len(cats)

    status_counts = Counter(
        (
            cat.status
            or "Не указан"
        )
        for cat in cats
    )

    mood_counts = Counter(
        (
            cat.mood
            or "Не указано"
        )
        for cat in cats
    )

    breed_counts = Counter(
        (
            cat.breed
            or "Без породы"
        )
        for cat in cats
    )

    vaccinated_count = sum(
        1
        for cat in cats
        if cat.is_vaccinated is True
    )

    sterilized_count = sum(
        1
        for cat in cats
        if cat.is_sterilized is True
    )

    today = date.today()

    need_vacc = []

    for cat in cats:

        if (
            cat.is_vaccinated is not True
            or (
                cat.next_vacc_date
                and cat.next_vacc_date < today
            )
        ):
            need_vacc.append(
                cat
            )

    stats = {
        "total":
            total,

        "in_shelter":
            status_counts.get(
                "В приюте",
                0
            ),

        "adopted":
            status_counts.get(
                "Усыновлён",
                0
            ),

        "treatment":
            status_counts.get(
                "На лечении",
                0
            ),

        "quarantine":
            status_counts.get(
                "В карантине",
                0
            ),

        "vaccinated_count":
            vaccinated_count,

        "sterilized_count":
            sterilized_count,

        "need_vacc_count":
            len(
                need_vacc
            ),

        "vaccination_rate":
            (
                vaccinated_count
                / total
                * 100
            )
            if total
            else 0,

        "sterilization_rate":
            (
                sterilized_count
                / total
                * 100
            )
            if total
            else 0,

        "need_vacc_rate":
            (
                len(need_vacc)
                / total
                * 100
            )
            if total
            else 0,

        "status_counts":
            dict(
                status_counts
            ),

        "mood_counts":
            dict(
                mood_counts
            ),

        "breed_counts":
            dict(
                breed_counts
            ),
    }

    return stats


# ==========================================================
# EXCEL
# ==========================================================

def export_to_excel(cats):

    cats = list(cats)

    workbook = Workbook()

    sheet = workbook.active

    sheet.title = "Котики"

    headers = [
        "ID",
        "Имя",
        "Кличка",
        "Возраст",
        "Порода",
        "Пол",
        "Окрас",
        "Вес",
        "Настроение",
        "Статус",
        "Дата поступления",
        "Дата усыновления",
        "Вакцинирован",
        "Последняя вакцинация",
        "Следующая вакцинация",
        "Стерилизован",
        "Характер",
        "Активность",
        "Любимая игрушка",
        "Любимое место",
        "Отношение к детям",
        "Отношение к кошкам",
        "Отношение к собакам",
        "Голос",
        "Питание",
        "История",
        "Рекомендация",
    ]

    sheet.append(
        headers
    )

    pink_fill = PatternFill(
        fill_type="solid",
        fgColor="EFA7C5"
    )

    header_font = Font(
        bold=True,
        color="FFFFFF"
    )

    thin = Side(
        style="thin",
        color="E8D7DF"
    )

    border = Border(
        left=thin,
        right=thin,
        top=thin,
        bottom=thin
    )

    for cell in sheet[1]:

        cell.fill = pink_fill

        cell.font = header_font

        cell.alignment = Alignment(
            horizontal="center",
            vertical="center",
            wrap_text=True
        )

        cell.border = border

    def yes_no(value):

        if value is True:
            return "Да"

        if value is False:
            return "Нет"

        return ""

    for cat in cats:

        sheet.append(
            [
                cat.id,
                cat.name,
                cat.nickname or "",
                cat.age
                if cat.age is not None
                else "",
                cat.breed or "",
                cat.gender or "",
                cat.color or "",
                cat.weight
                if cat.weight is not None
                else "",
                cat.mood or "",
                cat.status or "",
                cat.arrival_date or "",
                cat.adoption_date or "",
                yes_no(
                    cat.is_vaccinated
                ),
                cat.last_vacc_date or "",
                cat.next_vacc_date or "",
                yes_no(
                    cat.is_sterilized
                ),
                cat.personality or "",
                cat.activity_level or "",
                cat.favorite_toy or "",
                cat.favorite_place or "",
                yes_no(
                    cat.kids_friendly
                ),
                yes_no(
                    cat.cats_friendly
                ),
                yes_no(
                    cat.dogs_friendly
                ),
                cat.vocal or "",
                cat.food_preferences or "",
                cat.story or "",
                cat.adoption_recommendation
                or "",
            ]
        )

    for row in sheet.iter_rows(
        min_row=2
    ):

        for cell in row:

            cell.alignment = Alignment(
                vertical="top",
                wrap_text=True
            )

            cell.border = border

    widths = {
        "A": 8,
        "B": 18,
        "C": 18,
        "D": 10,
        "E": 18,
        "F": 12,
        "G": 18,
        "H": 10,
        "I": 18,
        "J": 18,
        "K": 18,
        "L": 18,
        "M": 16,
        "N": 20,
        "O": 20,
        "P": 16,
        "Q": 32,
        "R": 18,
        "S": 24,
        "T": 24,
        "U": 18,
        "V": 18,
        "W": 18,
        "X": 18,
        "Y": 32,
        "Z": 40,
        "AA": 40,
    }

    for column, width in widths.items():

        sheet.column_dimensions[
            column
        ].width = width

    sheet.freeze_panes = "A2"

    sheet.auto_filter.ref = (
        sheet.dimensions
    )


    stats_sheet = (
        workbook.create_sheet(
            "Статистика"
        )
    )

    stats = calculate_stats(
        cats
    )

    stats_rows = [
        [
            "Показатель",
            "Значение"
        ],
        [
            "Всего котиков",
            stats["total"]
        ],
        [
            "В приюте",
            stats["in_shelter"]
        ],
        [
            "Усыновлено",
            stats["adopted"]
        ],
        [
            "На лечении",
            stats["treatment"]
        ],
        [
            "В карантине",
            stats["quarantine"]
        ],
        [
            "Вакцинировано",
            stats[
                "vaccinated_count"
            ]
        ],
        [
            "Стерилизовано",
            stats[
                "sterilized_count"
            ]
        ],
        [
            "Требуют внимания по вакцинации",
            stats[
                "need_vacc_count"
            ]
        ],
    ]

    for row in stats_rows:

        stats_sheet.append(
            row
        )

    for cell in stats_sheet[1]:

        cell.fill = pink_fill
        cell.font = header_font
        cell.border = border

    for row in stats_sheet.iter_rows():

        for cell in row:

            cell.border = border

            cell.alignment = Alignment(
                vertical="top",
                wrap_text=True
            )

    stats_sheet.column_dimensions[
        "A"
    ].width = 38

    stats_sheet.column_dimensions[
        "B"
    ].width = 18


    output = BytesIO()

    workbook.save(
        output
    )

    output.seek(
        0
    )

    return output


# ==========================================================
# PDF
# ==========================================================

def _register_pdf_fonts():

    regular_candidates = [
        "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf",
        "/usr/share/fonts/truetype/liberation2/LiberationSans-Regular.ttf",
        "C:/Windows/Fonts/arial.ttf",
        "C:/Windows/Fonts/calibri.ttf",
    ]

    bold_candidates = [
        "/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf",
        "/usr/share/fonts/truetype/liberation2/LiberationSans-Bold.ttf",
        "C:/Windows/Fonts/arialbd.ttf",
        "C:/Windows/Fonts/calibrib.ttf",
    ]

    regular_path = next(
        (
            path
            for path
            in regular_candidates
            if os.path.exists(
                path
            )
        ),
        None
    )

    bold_path = next(
        (
            path
            for path
            in bold_candidates
            if os.path.exists(
                path
            )
        ),
        None
    )

    if not regular_path:

        raise FileNotFoundError(
            "Не найден шрифт "
            "с поддержкой кириллицы"
        )

    if not bold_path:

        bold_path = regular_path

    try:

        pdfmetrics.registerFont(
            TTFont(
                "KotopolisRegular",
                regular_path
            )
        )

    except Exception:

        pass

    try:

        pdfmetrics.registerFont(
            TTFont(
                "KotopolisBold",
                bold_path
            )
        )

    except Exception:

        pass

    return (
        "KotopolisRegular",
        "KotopolisBold"
    )


def generate_pdf_report(
    cats,
    stats=None
):

    cats = list(
        cats
    )

    if stats is None:

        stats = calculate_stats(
            cats
        )

    regular_font, bold_font = (
        _register_pdf_fonts()
    )

    buffer = BytesIO()

    document = SimpleDocTemplate(
        buffer,
        pagesize=landscape(
            A4
        ),
        leftMargin=14 * mm,
        rightMargin=14 * mm,
        topMargin=14 * mm,
        bottomMargin=14 * mm,
        title="Отчёт Котополис",
        author="Котополис",
    )

    styles = getSampleStyleSheet()

    title_style = ParagraphStyle(
        "KotopolisTitle",
        parent=styles["Title"],
        fontName=bold_font,
        fontSize=20,
        leading=24,
        textColor=colors.HexColor(
            "#D85F93"
        ),
        spaceAfter=8 * mm,
    )

    heading_style = ParagraphStyle(
        "KotopolisHeading",
        parent=styles["Heading2"],
        fontName=bold_font,
        fontSize=13,
        leading=16,
        textColor=colors.HexColor(
            "#D85F93"
        ),
        spaceBefore=4 * mm,
        spaceAfter=3 * mm,
    )

    normal_style = ParagraphStyle(
        "KotopolisNormal",
        parent=styles["BodyText"],
        fontName=regular_font,
        fontSize=9,
        leading=12,
    )

    small_style = ParagraphStyle(
        "KotopolisSmall",
        parent=normal_style,
        fontSize=7.5,
        leading=9,
    )

    story = []

    story.append(
        Paragraph(
            "Котополис 🐱",
            title_style
        )
    )

    story.append(
        Paragraph(
            "Отчёт о состоянии приюта",
            heading_style
        )
    )

    story.append(
        Paragraph(
            (
                "Отчёт сформирован: "
                + datetime.now().strftime(
                    "%d.%m.%Y %H:%M"
                )
            ),
            normal_style
        )
    )

    story.append(
        Spacer(
            1,
            5 * mm
        )
    )


    stats_data = [
        [
            Paragraph(
                "Показатель",
                normal_style
            ),
            Paragraph(
                "Значение",
                normal_style
            ),
        ],
        [
            "Всего котиков",
            stats.get(
                "total",
                0
            ),
        ],
        [
            "В приюте",
            stats.get(
                "in_shelter",
                0
            ),
        ],
        [
            "Усыновлено",
            stats.get(
                "adopted",
                0
            ),
        ],
        [
            "На лечении",
            stats.get(
                "treatment",
                0
            ),
        ],
        [
            "В карантине",
            stats.get(
                "quarantine",
                0
            ),
        ],
        [
            "Вакцинировано",
            stats.get(
                "vaccinated_count",
                0
            ),
        ],
        [
            "Стерилизовано",
            stats.get(
                "sterilized_count",
                0
            ),
        ],
        [
            "Требуют внимания по вакцинации",
            stats.get(
                "need_vacc_count",
                0
            ),
        ],
    ]

    stats_table = Table(
        stats_data,
        colWidths=[
            90 * mm,
            45 * mm
        ]
    )

    stats_table.setStyle(
        TableStyle(
            [
                (
                    "FONTNAME",
                    (0, 0),
                    (-1, -1),
                    regular_font
                ),
                (
                    "FONTNAME",
                    (0, 0),
                    (-1, 0),
                    bold_font
                ),
                (
                    "BACKGROUND",
                    (0, 0),
                    (-1, 0),
                    colors.HexColor(
                        "#EFA7C5"
                    )
                ),
                (
                    "TEXTCOLOR",
                    (0, 0),
                    (-1, 0),
                    colors.white
                ),
                (
                    "GRID",
                    (0, 0),
                    (-1, -1),
                    0.5,
                    colors.HexColor(
                        "#E6CDD7"
                    )
                ),
                (
                    "VALIGN",
                    (0, 0),
                    (-1, -1),
                    "TOP"
                ),
                (
                    "LEFTPADDING",
                    (0, 0),
                    (-1, -1),
                    6
                ),
                (
                    "RIGHTPADDING",
                    (0, 0),
                    (-1, -1),
                    6
                ),
                (
                    "TOPPADDING",
                    (0, 0),
                    (-1, -1),
                    5
                ),
                (
                    "BOTTOMPADDING",
                    (0, 0),
                    (-1, -1),
                    5
                ),
            ]
        )
    )

    story.append(
        stats_table
    )

    story.append(
        PageBreak()
    )

    story.append(
        Paragraph(
            "Список котиков",
            heading_style
        )
    )


    cat_rows = [
        [
            Paragraph(
                "ID",
                small_style
            ),
            Paragraph(
                "Имя",
                small_style
            ),
            Paragraph(
                "Возраст",
                small_style
            ),
            Paragraph(
                "Порода",
                small_style
            ),
            Paragraph(
                "Пол",
                small_style
            ),
            Paragraph(
                "Статус",
                small_style
            ),
            Paragraph(
                "Поступление",
                small_style
            ),
            Paragraph(
                "Вакцина",
                small_style
            ),
            Paragraph(
                "Стерил.",
                small_style
            ),
        ]
    ]

    for cat in cats:

        cat_rows.append(
            [
                str(
                    cat.id
                ),
                Paragraph(
                    str(
                        _pdf_escape(cat.name)
                        or ""
                    ),
                    small_style
                ),
                str(
                    cat.age
                    if cat.age is not None
                    else ""
                ),
                Paragraph(
                    str(
                        _pdf_escape(cat.breed)
                        or ""
                    ),
                    small_style
                ),
                Paragraph(
                    str(
                        _pdf_escape(cat.gender)
                        or ""
                    ),
                    small_style
                ),
                Paragraph(
                    str(
                        _pdf_escape(cat.status)
                        or ""
                    ),
                    small_style
                ),
                (
                    cat.arrival_date.strftime(
                        "%d.%m.%Y"
                    )
                    if cat.arrival_date
                    else ""
                ),
                (
                    "Да"
                    if cat.is_vaccinated
                    else "Нет"
                ),
                (
                    "Да"
                    if cat.is_sterilized
                    else "Нет"
                ),
            ]
        )

    cats_table = Table(
        cat_rows,
        repeatRows=1,
        colWidths=[
            10 * mm,
            28 * mm,
            17 * mm,
            35 * mm,
            22 * mm,
            30 * mm,
            27 * mm,
            22 * mm,
            22 * mm,
        ]
    )

    cats_table.setStyle(
        TableStyle(
            [
                (
                    "FONTNAME",
                    (0, 0),
                    (-1, -1),
                    regular_font
                ),
                (
                    "FONTNAME",
                    (0, 0),
                    (-1, 0),
                    bold_font
                ),
                (
                    "BACKGROUND",
                    (0, 0),
                    (-1, 0),
                    colors.HexColor(
                        "#EFA7C5"
                    )
                ),
                (
                    "TEXTCOLOR",
                    (0, 0),
                    (-1, 0),
                    colors.white
                ),
                (
                    "GRID",
                    (0, 0),
                    (-1, -1),
                    0.35,
                    colors.HexColor(
                        "#E7D6DD"
                    )
                ),
                (
                    "VALIGN",
                    (0, 0),
                    (-1, -1),
                    "TOP"
                ),
                (
                    "ALIGN",
                    (0, 0),
                    (0, -1),
                    "CENTER"
                ),
                (
                    "ROWBACKGROUNDS",
                    (0, 1),
                    (-1, -1),
                    [
                        colors.white,
                        colors.HexColor(
                            "#FFF7FA"
                        )
                    ]
                ),
                (
                    "TOPPADDING",
                    (0, 0),
                    (-1, -1),
                    4
                ),
                (
                    "BOTTOMPADDING",
                    (0, 0),
                    (-1, -1),
                    4
                ),
            ]
        )
    )

    story.append(
        cats_table
    )

    document.build(
        story
    )

    buffer.seek(
        0
    )

    return buffer
